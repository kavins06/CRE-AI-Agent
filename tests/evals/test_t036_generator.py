"""Public developer fixtures only: these tests never read existing eval answers."""

from __future__ import annotations

import csv
import hashlib
import json
import os
from decimal import Decimal, localcontext
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from typer.testing import CliRunner

from cre_brain.cli import create_app
from cre_brain.domain import CalcResult
from cre_brain.finance.proforma import build_proforma
from cre_brain.finance.rentroll import normalize_rent_roll
from cre_brain.finance.t12 import normalize_t12


def generator():
    from cre_brain.evals.commands import load_generator

    return load_generator()


def tree_bytes(root: Path) -> dict[str, bytes]:
    return {str(p.relative_to(root)): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def test_t036_ac1_generator_seeded_decimal_priors():
    gen = generator()
    priors = gen.DefaultPriors()
    assert priors.calibrated is False
    deals = [gen.sample_deal("public-test-seed", index=i) for i in range(1, 21)]
    assert deals == [gen.sample_deal("public-test-seed", index=i) for i in range(1, 21)]
    assert deals[0] != gen.sample_deal("different-public-seed", index=1)
    assert len({len(d.units) for d in deals}) > 1
    assert {d.layout for d in deals} == {"yardi", "realpage", "broker"}
    for deal in deals:
        assert priors.min_units <= len(deal.units) <= priors.max_units
        assert len({u.unit_id for u in deal.units}) == len(deal.units)
        assert isinstance(deal.proforma.base_monthly_revenue, Decimal)
        assert deal.proforma.projection_months == 12
        for unit in deal.units:
            assert isinstance(unit.market_rent, Decimal)
            assert 0 <= unit.monthly_concession <= unit.contract_rent <= unit.market_rent
            assert unit.occupied or unit.contract_rent == unit.monthly_concession == 0
        assert len({e.month for e in deal.entries}) == 12
        assert all(isinstance(e.amount, Decimal) and e.amount >= 0 for e in deal.entries)
    with localcontext() as ctx:
        ctx.prec = 5
        assert gen.sample_deal("public-test-seed", index=1) == deals[0]


@pytest.mark.parametrize(
    "update",
    [
        {"min_units": 0},
        {"min_units": 100, "max_units": 10},
        {"min_rent": Decimal("NaN")},
        {"min_rent": 1000.1},
        {"max_loss_to_lease_bps": 10001},
        {"calibrated": True},
    ],
)
def test_t036_ac1_generator_invalid_priors(update):
    with pytest.raises(ValueError):
        generator().DefaultPriors(**update)


def test_t036_ac1_generator_rejects_invalid_latent_math():
    gen = generator()
    deal = gen.sample_deal("public-test-seed")
    for updates in (
        {"units": (deal.units[0], deal.units[0])},
        {"deal_id": "../escape"},
        {"entries": deal.entries[:-1]},
        {"proforma": deal.proforma.model_copy(update={"monthly_reserves": Decimal("-1")})},
    ):
        bad = deal.model_copy(update=updates)
        with pytest.raises(ValueError):
            gen.calculate_truth(bad)


@pytest.mark.parametrize("layout", ["yardi", "realpage", "broker"])
@pytest.mark.parametrize("extension", ["xlsx", "csv"])
def test_t036_ac2_generator_rentroll_layouts_roundtrip(tmp_path, layout, extension):
    gen = generator()
    deal = gen.sample_deal("public-render-seed")
    path = tmp_path / f"roll.{extension}"
    gen.render_rent_roll(deal, path, layout=layout)
    if extension == "csv":
        with path.open(newline="") as handle:
            rows = list(csv.reader(handle))
    else:
        book = load_workbook(path, data_only=False)
        assert book.sheetnames == ["Rent Roll"]
        assert all(sheet.sheet_state == "visible" for sheet in book)
        assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
        rows = [
            [str(v) if v is not None else "" for v in row]
            for row in book.active.iter_rows(values_only=True)
        ]
        book.close()
    headers, data = rows[2], rows[3:]
    assert len(data) == len(deal.units)
    assert "USD/month" in " ".join(headers)
    fields = gen.RENT_ROLL_COLUMNS[layout]
    for row, unit in zip(data, deal.units, strict=True):
        values = dict(zip(fields, row, strict=True))
        assert values["unit_id"] == unit.unit_id
        assert Decimal(values["market_rent"]) == unit.market_rent
        assert Decimal(values["contract_rent"]) == unit.contract_rent
        assert Decimal(values["monthly_concession"]) == unit.monthly_concession
        assert values["status"] == ("Occupied" if unit.occupied else "Vacant")
    with pytest.raises(ValueError, match="exist"):
        gen.render_rent_roll(deal, path, layout=layout)


def test_t036_ac2_generator_t12_and_pdf(tmp_path):
    gen = generator()
    deal = gen.sample_deal("public-render-seed")
    gen.render_t12(deal, tmp_path / "t12.xlsx")
    book = load_workbook(tmp_path / "t12.xlsx")
    rows = list(book.active.values)
    assert len(rows[2]) == 13
    assert [str(v) for v in rows[2][1:]] == sorted({e.month.isoformat() for e in deal.entries})
    account_rows = {r[0]: r[1:] for r in rows[3:]}
    for entry in deal.entries:
        month_index = rows[2][1:].index(entry.month.isoformat())
        assert Decimal(str(account_rows[entry.account][month_index])) == entry.amount
    assert "below NOI" in " ".join(str(v) for row in rows for v in row)
    book.close()
    gen.render_om(deal, tmp_path / "om.pdf")
    reader = PdfReader(tmp_path / "om.pdf")
    text = " ".join(page.extract_text() for page in reader.pages)
    assert "Synthetic" in text and deal.deal_id in text
    assert str(deal.proforma.monthly_reserves) in text
    assert "below NOI" in text and "USD/month" in text
    assert "seed" not in text.lower() and "truth" not in text.lower()


def test_t036_ac3_generator_cli_twenty_separate_deterministic_packages(tmp_path):
    generator()  # Missing implementation must fail before interpreting CLI results.
    runner = CliRunner()
    for suffix in ("a", "b"):
        result = runner.invoke(
            create_app(),
            [
                "evals",
                "gen",
                "--n",
                "20",
                "--seed",
                "S",
                "--out-packages",
                str(tmp_path / f"packages-{suffix}"),
                "--out-truth",
                str(tmp_path / f"truth-{suffix}"),
            ],
        )
        assert result.exit_code == 0, result.output
        packages = tmp_path / f"packages-{suffix}"
        truth = tmp_path / f"truth-{suffix}"
        assert len(list(packages.iterdir())) == 20
        assert len(list(truth.glob("deal-*.json"))) == 20
        assert json.loads((truth / "manifest.json").read_text())["seed"] == "S"
        for deal_dir in packages.iterdir():
            assert len(list(deal_dir.iterdir())) == 8
            assert {p.suffix for p in deal_dir.iterdir()} == {".csv", ".xlsx", ".pdf"}
    assert tree_bytes(tmp_path / "packages-a") == tree_bytes(tmp_path / "packages-b")
    assert tree_bytes(tmp_path / "truth-a") == tree_bytes(tmp_path / "truth-b")


def test_t036_ac3_generator_finance_only_truth_lineage_and_reserves(tmp_path, monkeypatch):
    gen = generator()
    calls = []
    for name in ("normalize_rent_roll", "normalize_t12", "build_proforma"):
        original = getattr(gen.truth, name)

        def spy(*args, _original=original, _name=name, **kwargs):
            calls.append(_name)
            return _original(*args, **kwargs)

        monkeypatch.setattr(gen.truth, name, spy)
    gen.generate(
        n=1, seed="public-truth-test", out_packages=tmp_path / "p", out_truth=tmp_path / "t"
    )
    assert calls == ["normalize_rent_roll", "normalize_t12", "build_proforma"]
    data = json.loads((tmp_path / "t/deal-0001.json").read_text())
    deal = gen.LatentDeal.model_validate_json(json.dumps(data["latent"]))
    calcs = {k: CalcResult.model_validate(v) for k, v in data["calculations"].items()}
    expected = {
        "rent_roll": normalize_rent_roll(
            list(deal.units),
            calc_id=calcs["rent_roll"].calc_id,
            code_version=calcs["rent_roll"].code_version,
        ),
        "t12": normalize_t12(
            list(deal.entries),
            chart=deal.chart,
            missing_month_rule=deal.missing_month_rule,
            calc_id=calcs["t12"].calc_id,
            code_version=calcs["t12"].code_version,
        ),
        "proforma": build_proforma(
            deal.proforma,
            calc_id=calcs["proforma"].calc_id,
            code_version=calcs["proforma"].code_version,
        ),
    }
    assert calcs == expected
    assert all(c.code_version.startswith("sha256:") for c in calcs.values())
    input_ids = set(data["inputs"])
    assert all(set(c.inputs.values()) <= input_ids for c in calcs.values())
    pf = calcs["proforma"].outputs
    assert pf["year:1:noi"] - pf["year:1:cash_flow"] == pf["year:1:reserves"] > 0
    for name, checksum in data["package_sha256"].items():
        assert (
            hashlib.sha256((tmp_path / "p/deal-0001" / name).read_bytes()).hexdigest() == checksum
        )


@pytest.mark.parametrize(
    "case", ["same", "nested", "reverse", "symlink", "traversal", "existing", "file"]
)
def test_t036_ac3_generator_rejects_unsafe_paths(tmp_path, case):
    gen = generator()
    packages, truth = tmp_path / "p", tmp_path / "t"
    if case == "same":
        truth = packages
    elif case == "nested":
        truth = packages / "t"
    elif case == "reverse":
        packages = truth / "p"
    elif case == "symlink":
        (tmp_path / "alias").symlink_to(tmp_path, target_is_directory=True)
        truth = tmp_path / "alias/t"
    elif case == "traversal":
        truth = tmp_path / "other/../t"
    elif case == "existing":
        packages.mkdir()
    elif case == "file":
        packages.write_text("preserve me")
    before = tree_bytes(tmp_path)
    with pytest.raises(ValueError):
        gen.generate(n=1, seed="S", out_packages=packages, out_truth=truth)
    assert tree_bytes(tmp_path) == before
    assert not (tmp_path / "t").exists()


def test_t036_ac3_generator_render_failure_has_no_partial_output(tmp_path, monkeypatch):
    gen = generator()

    def fail(*args, **kwargs):
        raise RuntimeError("injected renderer failure")

    monkeypatch.setattr(gen.batch, "render_om", fail)
    with pytest.raises(RuntimeError, match="injected"):
        gen.generate(n=2, seed="S", out_packages=tmp_path / "p", out_truth=tmp_path / "t")
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("n", [0, -1, True, 1001])
def test_t036_ac3_generator_invalid_count_no_output(tmp_path, n):
    with pytest.raises(ValueError):
        generator().generate(n=n, seed="S", out_packages=tmp_path / "p", out_truth=tmp_path / "t")
    assert list(tmp_path.iterdir()) == []


def test_t036_ac1_generator_rejects_stale_input_identity():
    gen = generator()
    deal = gen.sample_deal("public-identities")
    unit = deal.units[0].model_copy(
        update={"market_rent": deal.units[0].market_rent + Decimal("1")}
    )
    mutated = deal.model_copy(update={"units": (unit, *deal.units[1:])})
    with pytest.raises(ValueError, match="identity"):
        gen.calculate_truth(mutated)


def test_t036_ac3_generator_publication_failure_rolls_back_truth(tmp_path, monkeypatch):
    gen = generator()
    paths = gen.batch.Destination.publish.__globals__
    original = paths["_rename_new"]

    def fail_second(parent, source, destination):
        if destination == "p":
            raise OSError("injected publish failure")
        original(parent, source, destination)

    monkeypatch.setitem(paths, "_rename_new", fail_second)
    with pytest.raises(OSError, match="injected"):
        gen.generate(n=1, seed="S", out_packages=tmp_path / "p", out_truth=tmp_path / "t")
    assert list(tmp_path.iterdir()) == []


def test_t036_ac3_generator_racing_destination_preserved(tmp_path, monkeypatch):
    gen = generator()
    paths = gen.batch.Destination.publish.__globals__
    original = paths["_rename_new"]

    def race(parent, source, destination):
        if destination == "p":
            (tmp_path / "p").mkdir()
            (tmp_path / "p/other-owner.txt").write_text("preserve")
        original(parent, source, destination)

    monkeypatch.setitem(paths, "_rename_new", race)
    with pytest.raises(ValueError, match="exist"):
        gen.generate(n=1, seed="S", out_packages=tmp_path / "p", out_truth=tmp_path / "t")
    assert tree_bytes(tmp_path) == {"p/other-owner.txt": b"preserve"}
    assert list(tmp_path.iterdir()) == [tmp_path / "p"]


def test_t036_ac3_generator_rejects_physical_parent_alias(tmp_path, monkeypatch):
    gen = generator()
    paths = gen.batch.Destination.__init__.__globals__
    original = paths["_open_directory"]
    (tmp_path / "alias").mkdir()
    monkeypatch.setitem(paths, "_open_directory", lambda path: original(tmp_path))
    with pytest.raises(ValueError, match="alias"):
        gen.generate(n=1, seed="S", out_packages=tmp_path / "p", out_truth=tmp_path / "alias/p")
    assert list(tmp_path.iterdir()) == [tmp_path / "alias"]


def test_t036_ac2_generator_failed_document_write_is_atomic(tmp_path, monkeypatch):
    gen = generator()
    from unittest.mock import Mock

    original = os.fdopen
    handle = Mock()
    handle.__enter__ = Mock(return_value=handle)
    handle.__exit__ = Mock(return_value=False)
    handle.write.side_effect = OSError("injected document write failure")

    def fail_write(fd, *args, **kwargs):
        # Close the real descriptor as well; this double doesn't leak an OS resource.
        original(fd, *args, **kwargs).close()
        return handle

    monkeypatch.setattr(os, "fdopen", fail_write)
    with pytest.raises(OSError, match="injected"):
        gen.render_om(gen.sample_deal("S"), tmp_path / "om.pdf")
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("layout", ["yardi", "realpage", "broker"])
def test_t036_ac2_generator_native_extraction_preserves_decimal_lexemes(layout):
    gen = generator()
    from cre_brain.extraction.preparse.models import Limits
    from cre_brain.extraction.preparse.native import xlsx_tables

    deal = gen.sample_deal("public-native-test")
    content = gen.render.rent_roll_bytes(deal, layout=layout, extension="xlsx")
    tables, external, hidden = xlsx_tables(content, "roll", "public-file-identity", Limits())
    assert not external and not hidden
    texts = [c.text for table in tables for c in table.cells]
    assert all(format(u.market_rent, "f") in texts for u in deal.units)
    assert all(c.kind != "formula" for table in tables for c in table.cells)
