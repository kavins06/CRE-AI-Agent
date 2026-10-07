"""Public synthetic deterministic core tests; no release or analyst-quality claim."""

import hashlib
from datetime import UTC, datetime
from decimal import Decimal, localcontext

import pytest
from pydantic import ValidationError
from sqlalchemy import create_engine

from cre_brain.config import load
from cre_brain.deliverables.uw import UWCore, UWModel, UWRefusal, UWRequest
from cre_brain.domain import CalcResult, ClaimType, Fact
from cre_brain.domain.base import TenantScope
from cre_brain.domain.models import Provenance
from cre_brain.excel.build_template import build_template
from cre_brain.runner.policy import HostContext
from cre_brain.runner.tools.contracts import FactAnchor, Reference, Template
from cre_brain.runner.tools.registry import ToolRegistry
from cre_brain.runner.tools.state import ToolState
from cre_brain.state.schema import metadata


class Inputs:
    def __init__(self):
        self.facts = {}
        self.templates = {}

    def fact(self, context, identity):
        return self.facts.get(identity)

    def template(self, context, identity):
        return self.templates.get(identity)

    def artifact(self, context, identity):
        return None

    def rule_policy(self, context, table):
        return None


@pytest.fixture
def uw(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'state.db'}")
    metadata.create_all(engine)
    inputs = Inputs()
    context = HostContext(
        scope=TenantScope(user_id="synthetic-user", firm_id="synthetic-firm"),
        task_id="uw-task",
        deal_id="synthetic-deal",
        role="lead",
        session_id="uw-session",
        release_id="synthetic-release",
        started_at=datetime.now(UTC),
    )
    registry = ToolRegistry(
        engine=engine,
        context=context,
        workspace=tmp_path,
        inputs=inputs,
        settings=load(apply_environment=False),
    )
    yield UWCore(registry), registry, inputs
    engine.dispose()


def put(uw, key, value, unit):
    _, registry, inputs = uw
    fact = Fact(
        fact_id=key,
        deal_id=registry.context.deal_id,
        key=key,
        value=value if type(value) is bool else Decimal(value),
        unit=unit,
        claim_type=ClaimType.VERIFIED_FACT,
        provenance=[Provenance(doc_id="public-synthetic", page=1)],
        known_at=datetime.now(UTC),
        version=1,
    )
    inputs.facts[key] = FactAnchor(fact=fact, authority="verified_source")
    result = registry.call("facts_put", {"anchor_id": key})
    assert result["status"] == "ok", result
    return Reference(kind="fact", record_id=key, version=1, key=key, unit=unit)


def request(uw):
    values = {
        "base_monthly_revenue": ("100000", "USD/month"),
        "base_monthly_operating_expenses": ("30000", "USD/month"),
        "annual_revenue_growth": ("0.03", "ratio"),
        "annual_expense_growth": ("0.02", "ratio"),
        "vacancy_rate": ("0.05", "ratio"),
        "credit_loss_rate": ("0.01", "ratio"),
        "monthly_reserves": ("2000", "USD/month"),
        "projection_months": ("60", "count"),
    }
    return UWRequest(proforma={k: put(uw, k, *v) for k, v in values.items()})


def component(uw, prefix, values):
    return {key: put(uw, f"{prefix}-{key}", value, unit) for key, (value, unit) in values.items()}


def test_t039_ac1_python_model_persisted_immutable_provenance(uw):
    core, registry, _ = uw
    req = request(uw)
    model = core.compose(req)
    assert isinstance(model, UWModel), model
    calc = model.calculation("proforma")
    assert calc.output("month:1:noi").value == Decimal("64050")
    assert calc.output("month:1:cash_flow").value == Decimal("62050")
    with registry.engine.connect() as connection:
        state = ToolState(connection, registry.context)
        stored = state.current(CalcResult, calc.calc_id)
        assert stored.outputs["month:1:noi"] == Decimal("64050")
        assert stored.code_version == calc.code_version
        assert len(stored.code_version) == 64
        assert set(stored.inputs.values()) == {r.record_id for r in req.proforma.values()}
    assert all(item.record_json and item.binding_json for item in model.evidence)
    with pytest.raises(ValidationError):
        model.calculations = ()
    assert core.compose(req) == model
    with localcontext() as context:
        context.prec = 3
        assert core.compose(req) == model


def test_t039_ac1_assumption_ranges_rationale_proxy_and_lineage(uw):
    core, registry, _ = uw
    req = request(uw)
    low = put(uw, "growth-low", "0.01", "ratio")
    high = put(uw, "growth-high", "0.05", "ratio")
    response = registry.call(
        "assumption_set",
        {
            "key": "revenue-growth",
            "value": req.proforma["annual_revenue_growth"].model_dump(),
            "low": low.model_dump(),
            "high": high.model_dump(),
            "rationale": "Synthetic proxy range supplied by host.",
            "is_proxy": True,
        },
    )
    assert response["status"] == "ok", response
    assumption = Reference(
        kind="assumption",
        record_id=response["data"]["record_id"],
        version=1,
        key="revenue-growth",
        unit="ratio",
    )
    req = req.model_copy(update={"proforma": {**req.proforma, "annual_revenue_growth": assumption}})
    model = core.compose(req)
    assert isinstance(model, UWModel), model
    evidence = next(e for e in model.evidence if e.reference == assumption)
    from cre_brain.domain import Assumption

    stored = Assumption.model_validate_json(evidence.record_json)
    assert (stored.low, stored.high, stored.is_proxy) == (Decimal("0.01"), Decimal("0.05"), True)
    assert stored.rationale == "Synthetic proxy range supplied by host."
    assert core.workbook(model).code == "unsupported_workbook"


def test_t039_ac1_missing_input_refuses_without_calculations(uw):
    core, registry, _ = uw
    req = request(uw)
    missing = req.model_copy(
        update={"proforma": {k: v for k, v in req.proforma.items() if k != "monthly_reserves"}}
    )
    result = core.compose(missing)
    assert isinstance(result, UWRefusal)
    assert result.code == "missing_inputs"
    assert "monthly_reserves" in result.questions[0].field
    with registry.engine.connect() as connection:
        assert ToolState(connection, registry.context).all_current(CalcResult) == []


def test_t039_ac2_tax_valueadd_debt_below_noi_and_canonical_link(uw):
    core, registry, _ = uw
    req = request(uw)
    taxes = component(
        uw,
        "tax",
        {
            "sale_price": ("10000000", "USD"),
            "current_assessed_value": ("5000000", "USD"),
            "current_annual_tax": ("60000", "USD"),
            "assessment_ratio": ("1", "ratio"),
            "millage": ("12", "mills"),
            "phase_in_years": ("1", "count"),
            "full_reassessment": (True, "bool"),
            "operating_expenses_exclude_property_tax": (True, "bool"),
        },
    )
    taxes["annual_growth_cap"] = None
    value_add = component(
        uw,
        "va",
        {
            "total_units": ("4", "count"),
            "units_per_month": ("2", "count"),
            "start_month": ("1", "count"),
            "downtime_months": ("1", "count"),
            "cost_per_unit": ("10000", "USD/unit"),
            "current_monthly_rent_per_unit": ("1000", "USD/unit/month"),
            "monthly_rent_premium": ("200", "USD/unit/month"),
            "ramp_months": ("1", "count"),
        },
    )
    debt = component(
        uw,
        "debt",
        {
            "property_value": ("10000000", "USD"),
            "max_ltv": ("0.7", "ratio"),
            "min_dscr": ("1.25", "ratio"),
            "min_debt_yield": ("0.08", "ratio"),
            "base_rate": ("0.04", "ratio"),
            "spread": ("0.01", "ratio"),
            "term_months": ("60", "count"),
            "amortization_months": ("360", "count"),
            "io_months": ("0", "count"),
            "fee_rate": ("0", "ratio"),
            "fixed_fee": ("0", "USD"),
            "prepayment_rate": ("0", "ratio"),
        },
    )
    model = core.compose(
        req.model_copy(update={"taxes": taxes, "value_add": value_add, "debt": debt})
    )
    assert isinstance(model, UWModel), model
    proforma = model.calculation("proforma")
    assert proforma.output("month:1:noi").value == Decimal("52169")
    assert proforma.output("month:1:cash_flow").value == Decimal("30169")
    assert model.calculation("taxes").output("target_annual_tax").value == Decimal("120000")
    assert model.calculation("value_add").output("total_renovation_cost").value == Decimal("40000")
    debt_calc = model.calculation("debt")
    amount = debt_calc.output("loan_amount").value
    assert amount == min(
        debt_calc.output(k).value for k in ("ltv_limit", "dscr_limit", "debt_yield_limit")
    )
    with registry.engine.connect() as connection:
        stored = ToolState(connection, registry.context).current(CalcResult, debt_calc.calc_id)
        assert proforma.calc_id in stored.inputs.values()
    assert core.workbook(model).code == "unsupported_workbook"


def test_uw_stale_and_unbound_evidence_refuses(uw):
    core, registry, _ = uw
    req = request(uw)
    forged = req.proforma["monthly_reserves"].model_copy(update={"version": 2})
    result = core.compose(
        req.model_copy(update={"proforma": {**req.proforma, "monthly_reserves": forged}})
    )
    assert isinstance(result, UWRefusal) and result.code == "stale_evidence"
    with registry.transaction() as state:
        state.event("stale", {"item_id": req.proforma["monthly_reserves"].record_id})
    assert core.compose(req).code == "stale_evidence"


@pytest.mark.parametrize(
    "flows,status", [(["-100", "230", "-132"], "ambiguous"), (["100", "200"], "undefined")]
)
def test_uw_undefined_ambiguous_irr_never_recommends(uw, flows, status):
    core, _, _ = uw
    req = request(uw)
    returns = {
        "cash_flows": [put(uw, f"flow-{i}", value, "USD") for i, value in enumerate(flows)],
        "finance_rate": put(uw, "finance-rate", "0.05", "ratio"),
        "reinvest_rate": put(uw, "reinvest-rate", "0.05", "ratio"),
    }
    from cre_brain.runner.tools.contracts import FinanceRun

    model = core.compose(
        req.model_copy(update={"returns": FinanceRun(fn="calculate_returns", args=returns)})
    )
    assert isinstance(model, UWModel), model
    assert model.irr_status == status
    assert model.recommendation is None


def test_uw_mf_standard_build_real_formulas_missing_engine_refuses(uw, tmp_path):
    core, registry, inputs = uw
    req = request(uw)
    path, mapping = build_template(tmp_path / "template", gates=registry.settings.gates)
    inputs.templates["mf_standard"] = Template(
        path=path.resolve(), mapping=mapping, sha256=hashlib.sha256(path.read_bytes()).hexdigest()
    )
    model = core.compose(req)
    result = core.workbook(model)
    assert isinstance(result, UWRefusal) and result.code == "missing_provider"
    # Building independently uses the actual trusted T032 handler and existing mapping.
    response = registry.call(
        "excel_build",
        {"template": "mf_standard", "calculation": model.calculation("proforma").calc_id},
    )
    assert response["status"] == "ok", response
    with registry.engine.connect() as connection:
        binding = ToolState(connection, registry.context).binding(
            response["data"]["artifact_id"], "workbook"
        )
    from openpyxl import load_workbook

    workbook = load_workbook(binding["descriptor"]["path"])
    assert workbook["Monthly"]["H2"].value == "=E2-F2"
    assert workbook["Monthly"]["I2"].value == "=H2-G2"
    workbook.close()


@pytest.mark.integration
def test_t039_ac1_uw_real_mf_standard_recalc_parity(uw, tmp_path):
    core, registry, inputs = uw
    from cre_brain.excel.libreoffice import LibreOfficeEngine

    registry.excel_engine = LibreOfficeEngine.from_environment(gates=registry.settings.gates)
    path, mapping = build_template(tmp_path / "template", gates=registry.settings.gates)
    inputs.templates["mf_standard"] = Template(
        path=path.resolve(), mapping=mapping, sha256=hashlib.sha256(path.read_bytes()).hexdigest()
    )
    model = core.compose(request(uw))
    result = core.workbook(model)
    assert result.status == "recalculated", result
    from openpyxl import load_workbook

    live = load_workbook(result.path)
    cached = load_workbook(result.path, data_only=True)
    assert live["Monthly"]["H2"].value == "=E2-F2"
    assert cached["Monthly"]["H2"].value == 64050
    assert cached["Monthly"]["I2"].value == 62050
    live.close()
    cached.close()


def test_uw_snapshot_tamper_refuses_before_workbook_provider(uw):
    core, _, _ = uw
    model = core.compose(request(uw))
    calc = model.calculation("proforma")
    forged_number = calc.numbers[0].model_copy(update={"value": Decimal("999999")})
    forged_calc = calc.model_copy(update={"numbers": (forged_number, *calc.numbers[1:])})
    result = core.workbook(model.model_copy(update={"calculations": (forged_calc,)}))
    assert isinstance(result, UWRefusal) and result.code == "untrusted_evidence"


@pytest.mark.parametrize("kind", ["unbound", "no_provenance", "wrong_unit"])
def test_uw_host_evidence_regressions(uw, kind):
    core, registry, inputs = uw
    req = request(uw)
    ref = req.proforma["monthly_reserves"]
    if kind == "wrong_unit":
        replacement = ref.model_copy(update={"unit": "ratio"})
    else:
        fact = inputs.facts[ref.record_id].fact
        identity = "alternate-reserves"
        record = fact.model_copy(
            update={
                "fact_id": identity,
                "key": identity,
                "provenance": [] if kind == "no_provenance" else fact.provenance,
            }
        )
        with registry.transaction() as state:
            # Simulate host ingestion defect using a separate source identity/key.
            state.append(record)
            if kind == "no_provenance":
                from cre_brain.runner.tools.state import fingerprint

                state.event(
                    "tool_result",
                    {
                        "binding": "fact",
                        "identity": identity,
                        "deal": registry.context.deal_id,
                        "digest": fingerprint(record),
                        "authority": "verified_source",
                    },
                )
        replacement = ref.model_copy(update={"record_id": identity, "key": identity})
    result = core.compose(
        req.model_copy(update={"proforma": {**req.proforma, "monthly_reserves": replacement}})
    )
    assert isinstance(result, UWRefusal)
    assert result.code == (
        "incompatible_evidence" if kind == "wrong_unit" else "untrusted_evidence"
    )


def test_uw_partial_optional_inputs_rollback_no_defaults(uw):
    core, registry, _ = uw
    req = request(uw)
    result = core.compose(
        req.model_copy(update={"debt": {"property_value": put(uw, "value", "10000000", "USD")}})
    )
    assert isinstance(result, UWRefusal) and result.code == "missing_inputs"
    with registry.engine.connect() as connection:
        assert ToolState(connection, registry.context).all_current(CalcResult) == []


def test_uw_reference_boundary_rejects_numbers_and_caller_pass(uw):
    req = request(uw)
    with pytest.raises(ValidationError):
        UWRequest.model_validate({"proforma": {"monthly_reserves": "1000"}})
    with pytest.raises(ValidationError):
        UWRequest.model_validate({**req.model_dump(), "passed": True})
    with pytest.raises(ValidationError):
        UWRequest.model_validate({**req.model_dump(), "gates": ["parity"]})


def test_uw_unsupported_renovation_cohort_count_refuses_before_execution(uw, monkeypatch):
    core, registry, _ = uw
    req = request(uw)
    plan = component(
        uw,
        "large-va",
        {
            "total_units": ("1000000", "count"),
            "units_per_month": ("1", "count"),
            "start_month": ("1", "count"),
            "downtime_months": ("0", "count"),
            "cost_per_unit": ("1000", "USD/unit"),
            "current_monthly_rent_per_unit": ("0", "USD/unit/month"),
            "monthly_rent_premium": ("100", "USD/unit/month"),
            "ramp_months": ("1", "count"),
        },
    )

    def must_not_execute(*args, **kwargs):
        pytest.fail("Unsupported unbounded renovation cohort loop was executed")

    monkeypatch.setattr(
        "cre_brain.deliverables.uw.calculations.build_value_add_schedule", must_not_execute
    )
    result = core.compose(req.model_copy(update={"value_add": plan}))
    assert isinstance(result, UWRefusal) and result.code == "unsupported_inputs"
    with registry.engine.connect() as connection:
        assert ToolState(connection, registry.context).all_current(CalcResult) == []


def test_uw_tax_uncapped_rule_has_explicit_persisted_manifest(uw):
    core, registry, _ = uw
    req = request(uw)
    taxes = component(
        uw,
        "uncapped-tax",
        {
            "sale_price": ("10000000", "USD"),
            "current_assessed_value": ("5000000", "USD"),
            "current_annual_tax": ("60000", "USD"),
            "assessment_ratio": ("1", "ratio"),
            "millage": ("12", "mills"),
            "phase_in_years": ("1", "count"),
            "full_reassessment": (True, "bool"),
            "operating_expenses_exclude_property_tax": (True, "bool"),
        },
    )
    taxes["annual_growth_cap"] = None
    model = core.compose(req.model_copy(update={"taxes": taxes}))
    assert isinstance(model, UWModel), model
    with registry.engine.connect() as connection:
        binding = ToolState(connection, registry.context).binding(
            model.calculation("taxes").calc_id, "calc"
        )
    assert "annual_growth_cap" in binding["field_references"]
    assert binding["field_references"]["annual_growth_cap"] is None


def set_growth(uw, req, *, key="revenue-growth", value=None):
    _, registry, _ = uw
    low = put(uw, f"{key}-low", "0.01", "ratio")
    high = put(uw, f"{key}-high", "0.05", "ratio")
    result = registry.call(
        "assumption_set",
        {
            "key": key,
            "value": (value or req.proforma["annual_revenue_growth"]).model_dump(),
            "low": low.model_dump(),
            "high": high.model_dump(),
            "rationale": "Host sourced growth scenario.",
            "is_proxy": False,
        },
    )
    assert result["status"] == "ok", result
    return Reference(
        kind="assumption", record_id=result["data"]["record_id"], version=1, key=key, unit="ratio"
    )


def with_growth(req, ref):
    return req.model_copy(update={"proforma": {**req.proforma, "annual_revenue_growth": ref}})


def ledger(registry):
    from sqlalchemy import select

    from cre_brain.state.schema import edges
    from cre_brain.state.store import tenant_filter

    with registry.engine.connect() as connection:
        state = ToolState(connection, registry.context)
        history = state.history(task_only=False)
        return (
            sum(e.kind == "usage" and e.payload.get("tools_call", False) for e in history),
            state.all_current(CalcResult),
            list(connection.execute(select(edges).where(tenant_filter(edges, state.scope)))),
            [e for e in history if e.kind == "tool_result" and e.payload.get("binding") == "calc"],
        )


def test_uw_review_superseded_assumption_refuses_latest_and_untouched_succeed(uw):
    core, registry, _ = uw
    req = request(uw)
    old = set_growth(uw, req)
    untouched = set_growth(uw, req, key="other-growth")
    latest = set_growth(uw, req, value=put(uw, "five-percent", "0.05", "ratio"))
    before = ledger(registry)
    result = core.compose(with_growth(req, old))
    assert isinstance(result, UWRefusal) and result.code == "stale_evidence"
    assert ledger(registry)[1:] == before[1:]
    model = core.compose(with_growth(req, latest))
    assert isinstance(model, UWModel), model
    assert model.calculation("proforma").output("month:13:noi").value == Decimal("68152.50")
    assert core.compose(with_growth(req, latest)) == model
    control = core.compose(with_growth(req, untouched))
    assert isinstance(control, UWModel), control
    assert control.calculation("proforma").output("month:13:noi").value == Decimal("66271.50")


@pytest.mark.parametrize("mismatch", ["deal", "key", "version", "user", "firm"])
def test_uw_review_assumption_scope_key_version_refuses(uw, mismatch):
    core, registry, _ = uw
    req = request(uw)
    ref = set_growth(uw, req)
    if mismatch == "deal":
        core = UWCore(
            registry.clone(
                context=registry.context.model_copy(
                    update={
                        "deal_id": "other",
                        "session_id": "other-session",
                        "task_id": "other-task",
                    }
                )
            )
        )
    elif mismatch in {"user", "firm"}:
        scope = registry.context.scope.model_copy(update={f"{mismatch}_id": "other"})
        core = UWCore(registry.clone(context=registry.context.model_copy(update={"scope": scope})))
    else:
        ref = ref.model_copy(update={"key": "other"} if mismatch == "key" else {"version": 2})
    result = core.compose(with_growth(req, ref))
    assert isinstance(result, UWRefusal)
    assert result.code in {"stale_evidence", "incompatible_evidence", "untrusted_evidence"}
    assert ledger(registry)[1] == []


def test_uw_review_missing_input_refusals_charge_until_cap(uw):
    from cre_brain.runner.policy import Limits

    _, registry, _ = uw
    core = UWCore(registry.clone(limits=Limits(max_tool_calls=11, max_session_calls=11)))
    req = request((core, core.registry, uw[2]))
    missing = req.model_copy(
        update={"proforma": {k: v for k, v in req.proforma.items() if k != "monthly_reserves"}}
    )
    assert [core.compose(missing).code for _ in range(4)] == [
        "missing_inputs",
        "missing_inputs",
        "missing_inputs",
        "budget_exceeded",
    ]
    assert ledger(registry) == (11, [], [], [])


@pytest.mark.parametrize("failure", ["typed", "validation", "arithmetic", "invalid_boundary"])
def test_uw_review_failure_charges_and_rolls_back_writes(uw, monkeypatch, failure):
    import cre_brain.deliverables.uw.core as core_module
    from cre_brain.runner.policy import Refusal

    core, registry, _ = uw
    req = request(uw)
    original = core_module.compose

    def fail_after_write(state, request):
        original(state, request)
        if failure == "typed":
            raise Refusal("missing_inputs", "Host needs another input.")
        if failure == "arithmetic":
            raise ArithmeticError("Unsupported arithmetic")
        UWRequest.model_validate({"proforma": {"bad": "literal"}})

    if failure == "invalid_boundary":
        req = req.model_copy(update={"proforma": {"bad": "literal"}})
    else:
        monkeypatch.setattr(core_module, "compose", fail_after_write)
    before = ledger(registry)
    result = core.compose(req)
    assert isinstance(result, UWRefusal)
    assert result.code == ("missing_inputs" if failure == "typed" else "invalid_input")
    assert ledger(registry) == (before[0] + 1, *before[1:])


def test_uw_review_deadline_after_snapshot_rolls_back_all_writes(uw, monkeypatch):
    from datetime import timedelta

    import cre_brain.deliverables.uw.core as core_module
    import cre_brain.runner.tools.registry as registry_module
    from cre_brain.runner.policy import Limits

    _, registry, _ = uw
    core = UWCore(registry.clone(limits=Limits(max_session_s=5)))
    req = request((core, core.registry, uw[2]))
    clock = [datetime.now(UTC)]

    class Clock:
        @staticmethod
        def now(tz):
            return clock[0]

    monkeypatch.setattr(registry_module, "datetime", Clock)
    original = core_module.compose

    def finish_after_deadline(state, request):
        model = original(state, request)
        clock[0] += timedelta(seconds=10)
        return model

    monkeypatch.setattr(core_module, "compose", finish_after_deadline)
    before = ledger(registry)
    result = core.compose(req)
    assert isinstance(result, UWRefusal) and result.code == "budget_exceeded"
    assert ledger(registry) == (before[0] + 1, *before[1:])


@pytest.mark.parametrize(
    "shape", ["extra", "oversized", "flows", "dates", "scalar", "fn", "partial_va"]
)
def test_uw_review_shape_refuses_before_any_resolution(uw, monkeypatch, shape):
    from cre_brain.runner.tools.contracts import FinanceRun

    core, _, _ = uw
    req = request(uw)
    ref = req.proforma["monthly_reserves"]
    returns = {"cash_flows": [ref, ref], "finance_rate": ref, "reinvest_rate": ref}
    if shape in {"extra", "oversized"}:
        fields = {f"extra-{i}": ref for i in range(1 if shape == "extra" else 1000)}
        req = req.model_copy(update={"proforma": {**req.proforma, **fields}})
    elif shape == "partial_va":
        req = req.model_copy(update={"value_add": {"total_units": ref}})
    else:
        if shape == "flows":
            returns["cash_flows"] = [ref] * 601
        elif shape == "dates":
            returns["dates"] = [ref] * 601
        elif shape == "scalar":
            returns["finance_rate"] = [ref]
        fn = "build_proforma" if shape == "fn" else "calculate_returns"
        req = req.model_copy(update={"returns": FinanceRun(fn=fn, args=returns)})

    def must_not_resolve(*args, **kwargs):
        pytest.fail("Shape refusal must precede source resolution")

    monkeypatch.setattr("cre_brain.deliverables.uw.evidence.resolve", must_not_resolve)
    result = core.compose(req)
    assert isinstance(result, UWRefusal)
    assert result.code in {"unsupported_inputs", "invalid_input", "missing_inputs"}


def test_uw_review_inspection_deduplicates_before_resolve(uw, monkeypatch):
    import cre_brain.deliverables.uw.evidence as evidence

    _, registry, _ = uw
    req = request(uw)
    ref = req.proforma["monthly_reserves"]
    original = evidence.resolve
    seen = []

    def count(state, ref):
        seen.append(ref)
        return original(state, ref)

    monkeypatch.setattr(evidence, "resolve", count)
    with registry.transaction() as state:
        snapshots = evidence.inspect_evidence(state, [ref] * 600)
    assert len(snapshots) == 1
    assert seen == [ref]


@pytest.mark.parametrize("months", ["0", "361", "1e32"])
def test_uw_review_projection_limit_before_finance_execution(uw, monkeypatch, months):
    core, registry, _ = uw
    req = request(uw)
    ref = put(uw, "huge-horizon", months, "count")
    req = req.model_copy(update={"proforma": {**req.proforma, "projection_months": ref}})

    def must_not_execute(*args, **kwargs):
        pytest.fail("Unsupported horizon reached a finance function")

    monkeypatch.setattr("cre_brain.deliverables.uw.calculations.run", must_not_execute)
    result = core.compose(req)
    assert isinstance(result, UWRefusal) and result.code == "invalid_input"
    assert ledger(registry)[1:] == ([], [], [])


@pytest.mark.parametrize("graph", ["fanout", "depth", "cycle", "expansion"])
def test_uw_review_dependency_bounds_before_recursive_resolution(uw, monkeypatch, graph):
    import cre_brain.deliverables.uw.evidence as evidence
    from cre_brain.runner.policy import Refusal
    from cre_brain.runner.tools.state import fingerprint

    _, registry, _ = uw
    root = request(uw).proforma["monthly_reserves"]
    with registry.transaction() as state:
        for index in range(21 if graph == "depth" else 5 if graph == "expansion" else 1):
            identity = f"bounded-calc-{index}"
            own = Reference(kind="calc", record_id=identity, version=1, key="result", unit="USD")
            sources = (
                [own]
                if graph == "cycle"
                else [root] * (10000 if graph == "fanout" else 16 if graph == "expansion" else 1)
            )
            calc = CalcResult(
                calc_id=identity,
                fn="build_proforma",
                inputs={"source": sources[0].record_id},
                outputs={"result": Decimal(1)},
                code_version=evidence.finance_version(),
            )
            state.append(calc)
            state.event(
                "tool_result",
                {
                    "binding": "calc",
                    "identity": identity,
                    "deal": state.context.deal_id,
                    "digest": fingerprint(calc),
                    "code_version": calc.code_version,
                    "units": {"result": "USD"},
                    "references": [r.model_dump(mode="json") for r in sources],
                },
            )
            root = own

    def must_not_resolve(*args, **kwargs):
        pytest.fail("Oversized/cyclic graph reached the shared recursive resolver")

    monkeypatch.setattr(evidence, "resolve", must_not_resolve)
    with registry.transaction() as state, pytest.raises(Refusal) as failure:
        evidence.inspect_evidence(state, [root])
    assert failure.value.category in {"invalid_input", "stale_evidence", "unsupported_inputs"}


@pytest.mark.parametrize(
    "field,value", [("start_month", "61"), ("ramp_months", "0"), ("units_per_month", "0")]
)
def test_uw_review_valueadd_counts_before_execution(uw, monkeypatch, field, value):
    core, registry, _ = uw
    req = request(uw)
    plan_values = {
        "total_units": ("4", "count"),
        "units_per_month": ("2", "count"),
        "start_month": ("1", "count"),
        "downtime_months": ("1", "count"),
        "cost_per_unit": ("10000", "USD/unit"),
        "current_monthly_rent_per_unit": ("1000", "USD/unit/month"),
        "monthly_rent_premium": ("200", "USD/unit/month"),
        "ramp_months": ("1", "count"),
    }
    plan_values[field] = (value, "count")
    plan = component(uw, "bounded-va", plan_values)

    def must_not_execute(*args, **kwargs):
        pytest.fail("Invalid counts reached value-add execution")

    monkeypatch.setattr(
        "cre_brain.deliverables.uw.calculations.build_value_add_schedule", must_not_execute
    )
    result = core.compose(req.model_copy(update={"value_add": plan}))
    assert isinstance(result, UWRefusal) and result.code == "invalid_input"
    assert ledger(registry)[1:] == ([], [], [])


@pytest.mark.parametrize("mismatch", ["deal", "key", "version", "user", "firm"])
def test_uw_review_direct_assumption_identity_boundary(uw, mismatch):
    from cre_brain.deliverables.uw.evidence import inspect_evidence
    from cre_brain.runner.policy import Refusal

    _, registry, _ = uw
    ref = set_growth(uw, request(uw))
    context = registry.context
    if mismatch == "deal":
        context = context.model_copy(update={"deal_id": "other-deal"})
    elif mismatch in {"user", "firm"}:
        context = context.model_copy(
            update={"scope": context.scope.model_copy(update={f"{mismatch}_id": "other"})}
        )
    else:
        ref = ref.model_copy(update={"key": "other-key"} if mismatch == "key" else {"version": 2})
    with registry.engine.connect() as connection, pytest.raises(Refusal):
        inspect_evidence(ToolState(connection, context), [ref])


def test_uw_review_maximum_property_projection_supported(uw):
    core, _, _ = uw
    req = request(uw)
    horizon = put(uw, "maximum-horizon", "360", "count")
    model = core.compose(
        req.model_copy(update={"proforma": {**req.proforma, "projection_months": horizon}})
    )
    assert isinstance(model, UWModel), model
    assert model.calculation("proforma").output("month:360:noi").value.is_finite()


def test_uw_review_large_property_counts_do_not_impose_unit_defaults(uw):
    core, _, _ = uw
    req = request(uw)
    plan = component(
        uw,
        "large-supported-va",
        {
            "total_units": ("60000", "count"),
            "units_per_month": ("100", "count"),
            "start_month": ("1", "count"),
            "downtime_months": ("600", "count"),
            "cost_per_unit": ("1000", "USD/unit"),
            "current_monthly_rent_per_unit": ("0", "USD/unit/month"),
            "monthly_rent_premium": ("100", "USD/unit/month"),
            "ramp_months": ("600", "count"),
        },
    )
    horizon = put(uw, "short-horizon", "1", "count")
    req = req.model_copy(
        update={"proforma": {**req.proforma, "projection_months": horizon}, "value_add": plan}
    )
    model = core.compose(req)
    assert isinstance(model, UWModel), model
    assert model.calculation("value_add").output("total_units").value == Decimal(60000)
    assert model.calculation("value_add").output("total_renovation_cost").value == Decimal(60000000)


def test_uw_review_nonconventional_returns_limit_before_execution(uw, monkeypatch):
    from cre_brain.runner.tools.contracts import FinanceRun

    core, registry, _ = uw
    req = request(uw)
    negative = put(uw, "alternating-negative", "-100", "USD")
    positive = put(uw, "alternating-positive", "100", "USD")
    rate = put(uw, "alternating-rate", "0.05", "ratio")
    returns = FinanceRun(
        fn="calculate_returns",
        args={
            "cash_flows": [negative if i % 2 else positive for i in range(129)],
            "finance_rate": rate,
            "reinvest_rate": rate,
        },
    )

    def must_not_execute(*args, **kwargs):
        pytest.fail("Unsupported nonconventional IRR reached the library")

    from cre_brain.runner.tools import finance

    schema, _, units = finance.SCHEMAS["calculate_returns"]
    monkeypatch.setitem(finance.SCHEMAS, "calculate_returns", (schema, must_not_execute, units))
    result = core.compose(req.model_copy(update={"returns": returns}))
    assert isinstance(result, UWRefusal) and result.code == "invalid_input"
    assert ledger(registry)[1:] == ([], [], [])


def test_uw_review_conventional_returns_600_flows_supported(uw):
    from cre_brain.runner.tools.contracts import FinanceRun

    core, _, _ = uw
    req = request(uw)
    flow = put(uw, "flow", "100", "USD")
    rate = put(uw, "rate", "0.05", "ratio")
    returns = FinanceRun(
        fn="calculate_returns",
        args={
            "cash_flows": [flow] * 600,
            "finance_rate": rate,
            "reinvest_rate": rate,
        },
    )
    model = core.compose(req.model_copy(update={"returns": returns}))
    assert isinstance(model, UWModel), model
    assert model.irr_status == "undefined"
    assert model.calculation("returns").output("total_distributions").value == Decimal(60000)
    assert core.compose(req.model_copy(update={"returns": returns})) == model


def test_uw_review_dated_returns_reuses_canonical_calculation(uw):
    from datetime import date

    from cre_brain.runner.tools.contracts import FinanceRun

    core, registry, inputs = uw
    req = request(uw)
    refs = []
    for index, day in enumerate((date(2023, 1, 1), date(2024, 1, 1))):
        identity = f"dated-flow-{index}"
        fact = Fact(
            fact_id=identity,
            deal_id=registry.context.deal_id,
            key=identity,
            value=day,
            unit="date",
            claim_type=ClaimType.VERIFIED_FACT,
            provenance=[Provenance(doc_id="public-synthetic", page=1)],
            known_at=datetime.now(UTC),
            version=1,
        )
        inputs.facts[identity] = FactAnchor(fact=fact, authority="verified_source")
        assert registry.call("facts_put", {"anchor_id": identity})["status"] == "ok"
        refs.append(
            Reference(kind="fact", record_id=identity, version=1, key=identity, unit="date")
        )
    rate = put(uw, "dated-rate", "0.05", "ratio")
    returns = FinanceRun(
        fn="calculate_returns",
        args={
            "cash_flows": [
                put(uw, "dated-contribution", "-100", "USD"),
                put(uw, "dated-distribution", "110", "USD"),
            ],
            "dates": refs,
            "finance_rate": rate,
            "reinvest_rate": rate,
        },
    )
    req = req.model_copy(update={"returns": returns})
    model = core.compose(req)
    assert isinstance(model, UWModel), model
    assert model.irr_status == "unique"
    assert abs(model.calculation("returns").output("xirr").value - Decimal("0.1")) < Decimal(
        "1e-24"
    )
    assert core.compose(req) == model


def test_uw_review_other_deal_same_key_does_not_supersede_assumption(uw):
    core, registry, inputs = uw
    req = request(uw)
    ref = set_growth(uw, req)
    other = registry.clone(
        context=registry.context.model_copy(
            update={
                "deal_id": "other-deal",
                "task_id": "other-task",
                "session_id": "other-session",
            }
        )
    )
    other_uw = (UWCore(other), other, inputs)
    value = put(other_uw, "other-value", "0.05", "ratio")
    low = put(other_uw, "other-low", "0.01", "ratio")
    high = put(other_uw, "other-high", "0.05", "ratio")
    result = other.call(
        "assumption_set",
        {
            "key": ref.key,
            "value": value.model_dump(),
            "low": low.model_dump(),
            "high": high.model_dump(),
            "rationale": "Independent other-deal scenario.",
            "is_proxy": False,
        },
    )
    assert result["status"] == "ok", result
    model = core.compose(with_growth(req, ref))
    assert isinstance(model, UWModel), model
    assert model.calculation("proforma").output("month:13:noi").value == Decimal("66271.50")
