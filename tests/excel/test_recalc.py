import xml.etree.ElementTree as ET
import zipfile
from decimal import Decimal, localcontext

import openpyxl
import pytest
from test_template import write

from cre_brain.config import load


def stale_workbook(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["A2"], sheet["A3"] = 1, 2, "=SUM(A1:A2)"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(files["xl/worksheets/sheet1.xml"])
    root.find(".//m:c[@r='A3']/m:v", ns).text = "999"
    files["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)


@pytest.fixture
def engine():
    from cre_brain.excel.libreoffice import LibreOfficeEngine

    return LibreOfficeEngine.from_environment(gates=load().gates)


@pytest.mark.integration
def test_t020_ac1_retains_formula_changed_cache_and_source(engine, tmp_path):
    path = tmp_path / "stale.xlsx"
    stale_workbook(path)
    assert openpyxl.load_workbook(path, data_only=True).active["A3"].value == 999
    output = engine.recalc(path)
    assert output != path
    assert openpyxl.load_workbook(output).active["A3"].value == "=SUM(A1:A2)"
    assert openpyxl.load_workbook(output, data_only=True).active["A3"].value == 3
    assert openpyxl.load_workbook(path, data_only=True).active["A3"].value == 999


def test_t020_ac2_exact_tolerances_and_missing_caches(tmp_path):
    from cre_brain.excel.parity import check_parity

    path = tmp_path / "values.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 1000000
    workbook.active["A2"] = 3
    workbook.save(path)
    gates = load().gates
    with localcontext() as context:
        context.prec = 2
        assert check_parity(path, {"Sheet!A1": Decimal("1000001")}, gates=gates).passed
        assert not check_parity(path, {"Sheet!A1": Decimal("1000001.00001")}, gates=gates).passed
    result = check_parity(path, {"Sheet!A3": Decimal("3")}, gates=gates)
    assert not result.passed and "Sheet!A3" in " ".join(result.failures)


@pytest.mark.parametrize("error", ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?"])
def test_t020_ac2_error_addresses(tmp_path, error):
    from cre_brain.excel.parity import check_parity

    workbook = openpyxl.Workbook()
    workbook.active["D7"] = error
    path = tmp_path / "errors.xlsx"
    workbook.save(path)
    result = check_parity(path, {"Sheet!A1": Decimal(1)}, gates=load().gates)
    assert not result.passed
    assert any("Sheet!D7" in f and error in f for f in result.failures)


@pytest.mark.integration
def test_t020_ac3_no_none_caches(engine, tmp_path):
    path = tmp_path / "stale.xlsx"
    stale_workbook(path)
    output = engine.recalc(path)
    assert openpyxl.load_workbook(output, data_only=True).active["A3"].value is not None


@pytest.mark.integration
def test_t020_ac4_synthetic_deal_pass_planted_error_fail(engine, template, synthetic, tmp_path):
    from cre_brain.excel.parity import check_parity

    result = write(template, synthetic, tmp_path / "deal.xlsx")
    output = engine.recalc(result.path)
    assert check_parity(output, result.expected, gates=load().gates).passed
    cached = openpyxl.load_workbook(output, data_only=True)
    assert all(cached[e.sheet][e.cell].value is not None for e in template[1].entries)
    workbook = openpyxl.load_workbook(result.path)
    workbook["Annual"]["G6"] = "=1/0"
    workbook.save(result.path)
    failed = check_parity(engine.recalc(result.path), result.expected, gates=load().gates)
    assert not failed.passed
    assert any("Annual!G6" in f and "#DIV/0!" in f for f in failed.failures)


def test_t020_ac4_graph_is_explicit_stub(tmp_path):
    from cre_brain.excel.graph import GraphExcelEngine

    with pytest.raises(NotImplementedError, match="licensed"):
        GraphExcelEngine().recalc(tmp_path / "deal.xlsx")


@pytest.mark.requires_license("MS_GRAPH")
def test_graph_separate_licensing_probe():
    from cre_brain.excel.graph import GraphExcelEngine

    assert GraphExcelEngine.implemented is False


@pytest.mark.parametrize("feature", ["macro", "external", "iterative", "data-table"])
def test_feature_validation_before_native_processing(tmp_path, feature):
    from cre_brain.excel.validation import validate_workbook

    path = tmp_path / "unsupported.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=SUM(B1:B2)"
    if feature == "iterative":
        workbook.calculation.iterate = True
    workbook.save(path)
    if feature == "macro":
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"not executed")
    elif feature == "external":
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    elif feature == "data-table":
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(files["xl/worksheets/sheet1.xml"])
        root.find(".//{*}f").set("t", "dataTable")
        files["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
    with pytest.raises(ValueError):
        validate_workbook(path, gates=load().gates)


def test_worker_rejects_missing_cache_in_saved_output(tmp_path, monkeypatch):
    """Subprocess boundary failure test; fake output is never recalc evidence."""
    import subprocess
    from pathlib import Path

    from cre_brain.excel.libreoffice import LibreOfficeConfig, LibreOfficeEngine

    path = tmp_path / "stale.xlsx"
    stale_workbook(path)
    observed = []

    class NoOpWorker:
        returncode = 0

        def __init__(self, args, **kwargs):
            observed.append((args, kwargs))
            workbook = openpyxl.Workbook()
            workbook.active["A3"] = "=SUM(A1:A2)"
            workbook.save(args[-1])

        def communicate(self, timeout):
            return None, b""

    monkeypatch.setattr(subprocess, "Popen", NoOpWorker)
    monkeypatch.setenv("SYNTHETIC_SECRET", "must-not-inherit")
    engine = LibreOfficeEngine(
        LibreOfficeConfig(soffice=Path("/usr/bin/python3"), python=Path("/usr/bin/python3")),
        gates=load().gates,
    )
    with pytest.raises(RuntimeError, match="missing recalculated cache"):
        engine.recalc(path)
    with pytest.raises(RuntimeError, match="missing recalculated cache"):
        engine.recalc(path)
    first, second = observed
    assert first[0][-3] != second[0][-3]
    for args, kwargs in observed:
        env = kwargs["env"]
        assert "SYNTHETIC_SECRET" not in env
        assert Path(env["HOME"]).parent == Path(args[-3])
        assert Path(env["TMPDIR"]).parent == Path(args[-3])
        assert not Path(args[-3]).exists()
    assert not list(tmp_path.glob("*.recalculated-*.xlsx"))


def test_worker_interrupt_cleans_owned_helper(tmp_path, monkeypatch):
    """Lifecycle-only failure test, no recalculation quality claim."""
    import subprocess
    from pathlib import Path

    from cre_brain.excel.libreoffice import LibreOfficeConfig, LibreOfficeEngine

    path = tmp_path / "stale.xlsx"
    stale_workbook(path)
    calls = []

    class InterruptedWorker:
        returncode = None

        def __init__(self, *args, **kwargs):
            pass

        def poll(self):
            return self.returncode

        def communicate(self, timeout):
            if not calls:
                raise KeyboardInterrupt()
            self.returncode = -15
            return None, b""

        def terminate(self):
            calls.append("owned-terminate")

        def kill(self):
            calls.append("owned-kill")

    monkeypatch.setattr(subprocess, "Popen", InterruptedWorker)
    engine = LibreOfficeEngine(
        LibreOfficeConfig(soffice=Path("/usr/bin/python3"), python=Path("/usr/bin/python3")),
        gates=load().gates,
    )
    with pytest.raises(KeyboardInterrupt):
        engine.recalc(path)
    assert calls == ["owned-terminate"]
