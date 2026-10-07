import json
from datetime import UTC, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine

from cre_brain.config import load
from cre_brain.domain import CalcResult, ClaimType, Fact, Provenance
from cre_brain.domain.base import TenantScope
from cre_brain.finance.proforma import ProFormaInput, build_proforma
from cre_brain.state.graph import DependencyGraph
from cre_brain.state.schema import metadata
from cre_brain.state.store import SqlVersionedStore

SCOPE = TenantScope(user_id="synthetic-user", firm_id="synthetic-firm")
VALUES = dict(
    base_monthly_revenue="100000",
    base_monthly_operating_expenses="40000",
    annual_revenue_growth=".03",
    annual_expense_growth=".02",
    vacancy_rate=".05",
    credit_loss_rate=".01",
    monthly_reserves="2000",
)


@pytest.fixture
def synthetic():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    metadata.create_all(engine)
    facts = SqlVersionedStore(engine, Fact)
    for key, value in VALUES.items():
        facts.append(
            Fact(
                fact_id=key,
                deal_id="synthetic",
                key=key,
                value=Decimal(value),
                unit="ratio" if "rate" in key or "growth" in key else "USD/month",
                claim_type=ClaimType.VERIFIED_FACT,
                provenance=[Provenance(doc_id="fixture")],
                known_at=datetime(2026, 10, 4, tzinfo=UTC),
                version=1,
            ),
            scope=SCOPE,
        )
    calc = build_proforma(
        ProFormaInput(
            input_id="base_monthly_revenue",
            projection_months=60,
            **{k: Decimal(v) for k, v in VALUES.items()},
        ),
        calc_id="stored-proforma",
        code_version="synthetic-v1",
    )
    calc = calc.model_copy(update={"inputs": {k: k for k in VALUES}})
    SqlVersionedStore(engine, CalcResult).append(calc, scope=SCOPE)
    yield engine, calc
    engine.dispose()


@pytest.fixture
def template(tmp_path):
    from cre_brain.excel.build_template import build_template

    return build_template(tmp_path, gates=load().gates)


def test_t019_ac1_template_whitelist_names_and_no_cycles(template):
    from cre_brain.excel.validation import validate_workbook

    path, mapping = template
    workbook = openpyxl.load_workbook(path)
    assert len(mapping.entries) > 300
    assert {e.name for e in mapping.entries} == set(workbook.defined_names)
    validate_workbook(path, gates=load().gates)
    formulas = [c.value for s in workbook for row in s for c in row if c.data_type == "f"]
    assert len(formulas) > 300
    assert all("INDIRECT" not in f for f in formulas)
    workbook["Inputs"]["B2"] = "=Inputs!B2"
    workbook.save(path)
    with pytest.raises(ValueError, match="[Cc]ircular"):
        validate_workbook(path, gates=load().gates)


def test_t019_ac2_template_map_roundtrip(template):
    from cre_brain.excel.models import TemplateMap

    path, mapping = template
    raw = json.loads(path.with_suffix(".map.json").read_text())
    assert TemplateMap.model_validate(raw) == mapping
    workbook = openpyxl.load_workbook(path)
    for entry in mapping.entries:
        assert list(workbook.defined_names[entry.name].destinations) == [(entry.sheet, entry.cell)]
    assert {e.key for e in mapping.entries if e.source == "fact"} == set(VALUES)
    assert any(e.key == "year:5:noi" and e.source == "calc" for e in mapping.entries)


def write(template, synthetic, target):
    from cre_brain.excel.writer import build_workbook

    engine, _ = synthetic
    return build_workbook(
        *template,
        target,
        engine=engine,
        scope=SCOPE,
        deal_id="synthetic",
        task_id="synthetic-task",
        calculations={"proforma": "stored-proforma"},
        gates=load().gates,
    )


def test_t019_ac3_template_writer_stored_identities(template, synthetic, tmp_path):
    from cre_brain.excel.validation import validate_workbook

    result = write(template, synthetic, tmp_path / "filled.xlsx")
    workbook = openpyxl.load_workbook(result.path)
    assert workbook["Inputs"]["B2"].value == 100000
    assert workbook["Monthly"]["B2"].data_type == "f"
    assert "base_monthly_revenue" in workbook["Inputs"]["B2"].comment.text
    assert "stored-proforma" in workbook["Monthly"]["B2"].comment.text
    assert result.expected["Annual!G6"] == synthetic[1].outputs["year:5:noi"]
    validate_workbook(result.path, gates=load().gates)


@pytest.mark.parametrize(
    "problem", ["missing", "ambiguous", "unit", "type", "conflict", "stale", "tenant", "lineage"]
)
def test_t019_ac3_template_rejects_incompatible_state(template, synthetic, tmp_path, problem):
    from cre_brain.excel.writer import build_workbook

    engine, calc = synthetic
    facts = SqlVersionedStore(engine, Fact)
    fact = facts.current("base_monthly_revenue", scope=SCOPE)
    scope = SCOPE
    if problem == "missing":
        calculations = {"proforma": "absent"}
    else:
        calculations = {"proforma": "stored-proforma"}
    if problem in {"unit", "type", "conflict"}:
        changes = (
            {"unit": "EUR/month"}
            if problem == "unit"
            else ({"value": "100000"} if problem == "type" else {"claim_type": ClaimType.CONFLICT})
        )
        facts.append(fact.model_copy(update={"version": 2, **changes}), scope=SCOPE)
    if problem == "ambiguous":
        facts.append(fact.model_copy(update={"fact_id": "duplicate"}), scope=SCOPE)
    if problem == "stale":
        graph = DependencyGraph(engine, release_id="fixture")
        graph.add_edge("base_monthly_revenue", "stored-proforma", scope=SCOPE)
        graph.mark_stale("base_monthly_revenue", task_id="synthetic-task", scope=SCOPE)
    if problem == "tenant":
        scope = TenantScope(user_id="other", firm_id=SCOPE.firm_id)
    if problem == "lineage":
        SqlVersionedStore(engine, CalcResult).append(
            calc.model_copy(update={"inputs": {"x": "absent"}}), scope=SCOPE
        )
    with pytest.raises(ValueError):
        build_workbook(
            *template,
            tmp_path / "bad.xlsx",
            engine=engine,
            scope=scope,
            deal_id="synthetic",
            task_id="synthetic-task",
            calculations=calculations,
            gates=load().gates,
        )
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "formula",
    [
        '=WEBSERVICE("https://invalid")',
        '=INDIRECT("A1")',
        "='[remote.xlsx]Sheet'!A1",
        "=MissingName",
        "=Monthly!B2",
    ],
)
def test_t019_ac1_template_rejects_unsupported_formulas(template, formula):
    from cre_brain.excel.validation import validate_workbook

    path, _ = template
    workbook = openpyxl.load_workbook(path)
    workbook["Monthly"]["B2"] = formula
    workbook.save(path)
    with pytest.raises(ValueError):
        validate_workbook(path, gates=load().gates)


@pytest.mark.parametrize(
    "problem", ["expired", "unvalidated-change", "precision", "future", "negative", "vacancy"]
)
def test_t019_ac3_template_rejects_stale_or_numerically_incompatible(
    template, synthetic, tmp_path, problem
):
    from datetime import date

    from cre_brain.domain import DateRange

    engine, _ = synthetic
    store = SqlVersionedStore(engine, Fact)
    key = "vacancy_rate" if problem == "vacancy" else "base_monthly_revenue"
    fact = store.current(key, scope=SCOPE)
    changes = {
        "expired": {"valid_time": DateRange(end=date(2000, 1, 1))},
        "future": {"valid_time": DateRange(start=date(2099, 1, 1))},
        "unvalidated-change": {"value": Decimal("200000")},
        "precision": {"value": Decimal("100000.00000000001")},
        "negative": {"value": Decimal("-1")},
        "vacancy": {"value": Decimal("1.1")},
    }[problem]
    store.append(fact.model_copy(update={"version": 2, **changes}), scope=SCOPE)
    with pytest.raises(ValueError):
        write(template, synthetic, tmp_path / "bad.xlsx")
    assert not (tmp_path / "bad.xlsx").exists()


def test_t019_ac1_template_reproducible_bytes(tmp_path, monkeypatch):
    import openpyxl.writer.excel as writer

    from cre_brain.excel.build_template import build_template

    class ClockA:
        @staticmethod
        def now(tz):
            return datetime(2026, 1, 1, tzinfo=tz)

    class ClockB:
        @staticmethod
        def now(tz):
            return datetime(2027, 1, 1, tzinfo=tz)

    from types import SimpleNamespace

    monkeypatch.setattr(
        writer,
        "datetime",
        SimpleNamespace(datetime=ClockA, timezone=__import__("datetime").timezone),
    )
    first, _ = build_template(tmp_path / "first", gates=load().gates)
    monkeypatch.setattr(
        writer,
        "datetime",
        SimpleNamespace(datetime=ClockB, timezone=__import__("datetime").timezone),
    )
    second, _ = build_template(tmp_path / "second", gates=load().gates)
    assert first.read_bytes() == second.read_bytes()
    assert (
        first.with_suffix(".map.json").read_bytes() == second.with_suffix(".map.json").read_bytes()
    )


def test_t019_ac3_template_unrelated_calculation_conflicts_do_not_block(
    template, synthetic, tmp_path
):
    engine, _ = synthetic
    store = SqlVersionedStore(engine, CalcResult)
    unrelated = CalcResult(
        calc_id="unrelated",
        fn="elsewhere",
        inputs={"x": "another-deal"},
        outputs={"x": Decimal(1)},
        code_version="v1",
    )
    store.append(unrelated, scope=SCOPE)
    store.append(unrelated.model_copy(update={"outputs": {"x": Decimal(2)}}), scope=SCOPE)
    assert write(template, synthetic, tmp_path / "filled.xlsx").path.is_file()
