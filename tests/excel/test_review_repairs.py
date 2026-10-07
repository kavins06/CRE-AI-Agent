"""Independent review regressions; subprocess doubles prove lifecycle only."""

import io
import signal
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import UTC, date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from test_recalc import stale_workbook
from test_template import SCOPE, write

from cre_brain.config import load
from cre_brain.domain import CalcResult, ClaimType, DateRange, Fact
from cre_brain.domain.base import TenantScope
from cre_brain.excel import _uno_worker
from cre_brain.excel.libreoffice import LibreOfficeConfig, LibreOfficeEngine
from cre_brain.excel.parity import check_parity
from cre_brain.excel.validation import validate_workbook
from cre_brain.excel.writer import build_workbook
from cre_brain.state.graph import DependencyGraph
from cre_brain.state.store import SqlVersionedStore


def native_engine():
    return LibreOfficeEngine(
        LibreOfficeConfig(soffice="/usr/bin/python3", python="/usr/bin/python3", timeout_s=1),
        gates=load().gates,
    )


def build(template, synthetic, target, **kwargs):
    return build_workbook(
        *template,
        target,
        engine=synthetic[0],
        scope=SCOPE,
        deal_id="synthetic",
        task_id=kwargs.pop("task_id", "synthetic-task"),
        calculations=kwargs.pop("calculations", {"proforma": "stored-proforma"}),
        gates=load().gates,
        as_of=date(2026, 10, 4),
        **kwargs,
    )


def cache_expected(path, expected):
    """Plant synthetic caches without flattening formulas; never UNO evidence."""
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(("Inputs", "Monthly", "Annual"), 1):
        name = f"xl/worksheets/sheet{index}.xml"
        root = ET.fromstring(files[name])
        for cell in root.findall(".//{*}c"):
            address = f"{sheet}!{cell.get('r')}"
            if address in expected and cell.find("{*}f") is not None:
                cell.find("{*}v").text = str(expected[address])
        files[name] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("subset", ["input-only", "omit-output", "extra"])
def test_review_1_mapping_requires_exact_expectations(template, synthetic, tmp_path, subset):
    result = write(template, synthetic, tmp_path / "build.xlsx")
    expected = dict(result.expected)
    if subset == "input-only":
        expected = {"Inputs!B2": expected["Inputs!B2"]}
    elif subset == "omit-output":
        del expected["Annual!G6"]
    else:
        expected["Inputs!Z99"] = Decimal(0)
    with pytest.raises(ValueError, match="coverage"):
        check_parity(result.path, expected, gates=load().gates, mapping=template[1])


@pytest.mark.parametrize("replacement", [123, "=123", "=SUM(Monthly!H50:H60)"])
def test_review_2_mapped_formula_integrity(template, synthetic, tmp_path, replacement):
    result = write(template, synthetic, tmp_path / "build.xlsx")
    workbook = openpyxl.load_workbook(result.path)
    workbook["Annual"]["G6"] = replacement
    workbook.save(result.path)
    with pytest.raises(ValueError, match="formula"):
        check_parity(result.path, result.expected, gates=load().gates, mapping=template[1])


def test_review_2_engine_rejects_changed_formula_after_worker(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    stale_workbook(path)

    class Worker:
        returncode = 0

        def __init__(self, args, **kwargs):
            stale_workbook(Path(args[-1]))
            with zipfile.ZipFile(args[-1]) as archive:
                files = {n: archive.read(n) for n in archive.namelist()}
            root = ET.fromstring(files["xl/worksheets/sheet1.xml"])
            root.find(".//{*}f").text = "999"
            files["xl/worksheets/sheet1.xml"] = ET.tostring(root)
            with zipfile.ZipFile(args[-1], "w") as archive:
                for name, data in files.items():
                    archive.writestr(name, data)

        def communicate(self, timeout):
            return None, b""

    monkeypatch.setattr(subprocess, "Popen", Worker)
    with pytest.raises(RuntimeError, match="formula"):
        native_engine().recalc(path)


@pytest.mark.parametrize("position", ["mapped", "upstream"])
@pytest.mark.parametrize("problem", ["known-future", "naive", "expired", "valid-future"])
def test_review_3_all_fact_time_boundaries(template, synthetic, tmp_path, position, problem):
    engine, calc = synthetic
    store = SqlVersionedStore(engine, Fact)
    fact = store.current("base_monthly_revenue", scope=SCOPE)
    changes = {
        "known-future": {"known_at": datetime(2026, 10, 5, tzinfo=UTC)},
        "naive": {"known_at": datetime(2026, 10, 4)},
        "expired": {"valid_time": DateRange(end=date(2000, 1, 1))},
        "valid-future": {"valid_time": DateRange(start=date(2099, 1, 1))},
    }[problem]
    if position == "mapped":
        store.append(fact.model_copy(update={"version": 2, **changes}), scope=SCOPE)
        calculations = {"proforma": calc.calc_id}
    else:
        store.append(
            fact.model_copy(update={"fact_id": "upstream", "key": "upstream", **changes}),
            scope=SCOPE,
        )
        SqlVersionedStore(engine, CalcResult).append(
            calc.model_copy(
                update={"calc_id": "fresh", "inputs": {**calc.inputs, "x": "upstream"}}
            ),
            scope=SCOPE,
        )
        calculations = {"proforma": "fresh"}
    with pytest.raises(ValueError, match="known|valid time"):
        build(template, synthetic, tmp_path / "bad.xlsx", calculations=calculations)
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "known",
    [
        datetime(2026, 10, 4, 23, 59, 59, 999999, tzinfo=UTC),
        datetime(2026, 10, 5, 1, 59, 59, 999999, tzinfo=timezone(timedelta(hours=2))),
    ],
)
def test_review_3_known_at_includes_utc_end_of_day(template, synthetic, tmp_path, known):
    store = SqlVersionedStore(synthetic[0], Fact)
    fact = store.current("base_monthly_revenue", scope=SCOPE)
    store.append(fact.model_copy(update={"known_at": known, "version": 2}), scope=SCOPE)
    assert build(template, synthetic, tmp_path / "build.xlsx").path.exists()


@pytest.mark.parametrize("kind", ["calc", "fact"])
def test_review_4_stale_identity_cannot_move_tasks(template, synthetic, tmp_path, kind):
    graph = DependencyGraph(synthetic[0], release_id="test")
    identity = "stored-proforma" if kind == "calc" else "base_monthly_revenue"
    graph.add_edge("changed", identity, scope=SCOPE)
    graph.mark_stale("changed", task_id="old-task", scope=SCOPE)
    with pytest.raises(ValueError, match="[Ss]tale"):
        build(template, synthetic, tmp_path / "bad.xlsx", task_id="another-task")


def test_review_4_global_stale_keeps_task_order_and_tenant_isolation(synthetic):
    graph = DependencyGraph(synthetic[0], release_id="test")
    graph.add_edge("changed", "stored-proforma", scope=SCOPE)
    graph.add_edge("stored-proforma", "deliverable", scope=SCOPE)
    graph.mark_stale("changed", task_id="old-task", scope=SCOPE)
    assert graph.stale_items("new-task", scope=SCOPE) == []
    assert graph.stale_items("old-task", scope=SCOPE) == ["stored-proforma", "deliverable"]
    assert set(graph.tenant_stale_items(scope=SCOPE)) == {"stored-proforma", "deliverable"}
    for other in (
        TenantScope(user_id="other", firm_id=SCOPE.firm_id),
        TenantScope(user_id=SCOPE.user_id, firm_id="other"),
    ):
        assert graph.tenant_stale_items(scope=other) == []


def test_review_5_unrelated_root_cannot_borrow_lineage(template, synthetic, tmp_path):
    path, mapping = template
    altered = mapping.model_copy(deep=True)
    altered.entries[-1] = altered.entries[-1].model_copy(update={"calculation": "second"})
    engine, calc = synthetic
    SqlVersionedStore(engine, CalcResult).append(
        calc.model_copy(update={"calc_id": "incomplete", "inputs": {"x": "monthly_reserves"}}),
        scope=SCOPE,
    )
    with pytest.raises(ValueError, match="namespace|lineage"):
        build(
            (path, altered),
            synthetic,
            tmp_path / "bad.xlsx",
            calculations={"proforma": calc.calc_id, "second": "incomplete"},
        )


def test_review_6_rejects_local_name_shadowing(template):
    path, _ = template
    workbook = openpyxl.load_workbook(path)
    workbook["Monthly"].defined_names.add(
        DefinedName("base_monthly_revenue", attr_text="'Monthly'!$B$2", localSheetId=1)
    )
    workbook.save(path)
    with pytest.raises(ValueError, match="local.*name|name.*local"):
        validate_workbook(path, gates=load().gates)


@pytest.mark.parametrize("feature", ["content-type", "vba-relationship", "macro-sheet"])
def test_review_7_rejects_macro_semantics_independent_of_filename(tmp_path, feature):
    path = tmp_path / "innocent.xlsx"
    stale_workbook(path)
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    if feature == "content-type":
        root = ET.fromstring(files["[Content_Types].xml"])
        for node in root:
            if node.get("PartName") == "/xl/workbook.xml":
                node.set("ContentType", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        files["[Content_Types].xml"] = ET.tostring(root)
    else:
        root = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {
                "Id": "rId999",
                "Target": "payload.bin",
                "Type": (
                    "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
                    if feature == "vba-relationship"
                    else "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet"
                ),
            },
        )
        files["xl/_rels/workbook.xml.rels"] = ET.tostring(root)
        files["xl/payload.bin"] = b"inert synthetic payload"
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="macro|Macro|VBA|vba"):
        validate_workbook(path, gates=load().gates)


@pytest.mark.parametrize("failure", ["early-exit", "timeout-graceful", "interrupted"])
def test_review_8_sweeps_group_even_when_helper_exits_promptly(tmp_path, monkeypatch, failure):
    path = tmp_path / "source.xlsx"
    stale_workbook(path)
    groups = []
    observed = {}

    class Worker:
        pid = 123456789
        returncode = None
        calls = 0

        def __init__(self, args, **kwargs):
            observed.update(kwargs)

        def communicate(self, timeout):
            self.calls += 1
            if self.calls == 1:
                if failure == "timeout-graceful":
                    raise subprocess.TimeoutExpired("worker", timeout)
                if failure == "interrupted":
                    self.returncode = -15
                    raise KeyboardInterrupt()
            self.returncode = 1
            return None, b"private native text"

        def terminate(self):
            self.returncode = -15

    monkeypatch.setattr(subprocess, "Popen", Worker)
    monkeypatch.setattr("os.killpg", lambda pid, sig: groups.append((pid, sig)))
    with pytest.raises(KeyboardInterrupt if failure == "interrupted" else RuntimeError):
        native_engine().recalc(path)
    assert observed["start_new_session"] is True
    assert groups == [(Worker.pid, signal.SIGKILL)]


def test_review_8_worker_installs_handlers_before_spawning_office(tmp_path, monkeypatch):
    installed = {}
    at_spawn = {}
    monkeypatch.setattr(
        _uno_worker.signal, "signal", lambda sig, handler: installed.update({sig: handler})
    )
    monkeypatch.setattr(sys, "argv", ["worker", "office", str(tmp_path), "in", "out"])

    def spawn(*args, **kwargs):
        at_spawn.update(installed)
        raise RuntimeError("synthetic spawn failed")

    monkeypatch.setattr(subprocess, "Popen", spawn)
    with pytest.raises(RuntimeError):
        _uno_worker.main()
    assert at_spawn.get(signal.SIGTERM) is _uno_worker._interrupted
    assert at_spawn.get(signal.SIGINT) is _uno_worker._interrupted


@pytest.mark.parametrize("failure", ["stderr", "spawn", "communication"])
def test_review_9_public_native_errors_are_allowlisted(tmp_path, monkeypatch, failure):
    path = tmp_path / "source.xlsx"
    stale_workbook(path)
    secret = "Bearer SYNTHETIC_TOKEN /private/workbook text"

    class Worker:
        pid = 123456789
        returncode = 1

        def __init__(self, *args, **kwargs):
            if failure == "spawn":
                raise OSError(secret)

        def communicate(self, timeout):
            if failure == "communication":
                raise OSError(secret)
            return None, secret.encode()

    monkeypatch.setattr(subprocess, "Popen", Worker)
    monkeypatch.setattr("os.killpg", lambda *args: None)
    with pytest.raises(RuntimeError) as caught:
        native_engine().recalc(path)
    assert str(caught.value) == "LibreOffice UNO recalculation failed"
    assert caught.value.__cause__ is None
    assert caught.value.__suppress_context__


def test_review_9_worker_never_forwards_native_stderr_or_exception(tmp_path, monkeypatch, capsys):
    secret = "Bearer SYNTHETIC_TOKEN /private/workbook text"

    class Office:
        returncode = 1
        stderr = io.BytesIO(secret.encode())

        def poll(self):
            return self.returncode

    monkeypatch.setattr(subprocess, "Popen", lambda *args, **kwargs: Office())
    monkeypatch.setattr(_uno_worker.signal, "signal", lambda *args: None)
    monkeypatch.setattr(sys, "argv", ["worker", "office", str(tmp_path), "in", "out"])

    def import_uno(name):
        raise RuntimeError(secret)

    monkeypatch.setattr(_uno_worker.importlib, "import_module", import_uno)
    with pytest.raises(RuntimeError) as caught:
        _uno_worker.main()
    assert secret not in str(caught.value)
    assert secret not in capsys.readouterr().err


@pytest.mark.parametrize("changed", ["expectations", "formula-before", "formula-after"])
def test_review_deliverable_path_is_bound_to_build(template, synthetic, tmp_path, changed):
    from cre_brain.excel.deliverable import recalc_deliverable

    result = write(template, synthetic, tmp_path / "build.xlsx")
    calls = []

    def change_formula(path):
        workbook = openpyxl.load_workbook(path)
        workbook["Annual"]["G6"] = float(result.expected["Annual!G6"])
        workbook.save(path)

    class Engine:
        def recalc(self, path):
            calls.append(path)
            if changed == "formula-after":
                change_formula(path)
            cache_expected(path, result.expected)
            return path

    if changed == "expectations":
        result = result.model_copy(update={"expected": {"Inputs!B2": result.expected["Inputs!B2"]}})
    if changed == "formula-before":
        change_formula(result.path)
    with pytest.raises(ValueError, match="coverage|formula"):
        recalc_deliverable(result, Engine(), gates=load().gates)
    assert bool(calls) == (changed == "formula-after")


def test_review_deliverable_complete_build_compares_every_cell(template, synthetic, tmp_path):
    from cre_brain.excel.deliverable import recalc_deliverable

    result = write(template, synthetic, tmp_path / "build.xlsx")

    class Engine:
        def recalc(self, path):
            cache_expected(path, result.expected)
            return path

    path, gate = recalc_deliverable(result, Engine(), gates=load().gates)
    assert path == result.path
    assert gate.passed
    assert gate.metrics["compared_cells"] == 517


def test_review_2_normalization_preserves_sheet_identity():
    from cre_brain.excel.validation import formula_signature

    assert formula_signature("='Cash$'!A1") != formula_signature("=Cash!A1")


@pytest.mark.parametrize(
    "formula",
    [
        "=sum('Monthly'!$H$50:$H$61)",
        "=SUM(MONTHLY!H50:H61)",
    ],
)
def test_review_2_only_harmless_formula_normalization(template, synthetic, tmp_path, formula):
    result = write(template, synthetic, tmp_path / "build.xlsx")
    workbook = openpyxl.load_workbook(result.path)
    workbook["Annual"]["G6"] = formula
    workbook.save(result.path)
    cache_expected(result.path, result.expected)
    assert check_parity(
        result.path, result.expected, gates=load().gates, mapping=template[1]
    ).passed


@pytest.mark.parametrize("problem", ["conflict", "deal", "stale", "expired"])
def test_review_3_deep_upstream_fact_checks(template, synthetic, tmp_path, problem):
    engine, calc = synthetic
    facts = SqlVersionedStore(engine, Fact)
    upstream = facts.current("base_monthly_revenue", scope=SCOPE).model_copy(
        update={"fact_id": "deep", "key": "deep"}
    )
    changes = {
        "conflict": {"claim_type": ClaimType.CONFLICT},
        "deal": {"deal_id": "other"},
        "stale": {},
        "expired": {"valid_time": DateRange(end=date(2000, 1, 1))},
    }[problem]
    facts.append(upstream.model_copy(update=changes), scope=SCOPE)
    calcs = SqlVersionedStore(engine, CalcResult)
    calcs.append(
        calc.model_copy(update={"calc_id": "bridge", "inputs": {"x": "deep"}}), scope=SCOPE
    )
    calcs.append(
        calc.model_copy(update={"calc_id": "fresh", "inputs": {**calc.inputs, "x": "bridge"}}),
        scope=SCOPE,
    )
    if problem == "stale":
        graph = DependencyGraph(engine, release_id="test")
        graph.add_edge("changed", "deep", scope=SCOPE)
        graph.mark_stale("changed", task_id="other-task", scope=SCOPE)
    with pytest.raises(ValueError, match="conflict|deal|[Ss]tale|valid time"):
        build(template, synthetic, tmp_path / "bad.xlsx", calculations={"proforma": "fresh"})


def test_review_4_regeneration_uses_new_calc_identity(template, synthetic, tmp_path):
    engine, calc = synthetic
    graph = DependencyGraph(engine, release_id="test")
    graph.add_edge("base_monthly_revenue", calc.calc_id, scope=SCOPE)
    graph.mark_stale("base_monthly_revenue", task_id="previous-task", scope=SCOPE)
    SqlVersionedStore(engine, CalcResult).append(
        calc.model_copy(update={"calc_id": "fresh"}),
        scope=SCOPE,
    )
    assert build(
        template, synthetic, tmp_path / "build.xlsx", calculations={"proforma": "fresh"}
    ).path.exists()
    assert calc.calc_id in graph.tenant_stale_items(scope=SCOPE)
