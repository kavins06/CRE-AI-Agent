"""Compare saved recalculated caches to stored authoritative Decimal outputs."""

from decimal import Decimal
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook

from cre_brain.config.settings import GateSettings
from cre_brain.domain import GateResult
from cre_brain.excel.models import TemplateMap
from cre_brain.excel.validation import validate_mapped_workbook

ERRORS = frozenset({"#REF!", "#DIV/0!", "#VALUE!", "#NAME?"})


def check_parity(
    path: Path,
    expected: dict[str, Decimal],
    *,
    gates: GateSettings,
    mapping: TemplateMap | None = None,
) -> GateResult:
    """Numeric subset comparison unless a complete validated map is supplied.

    Product deliverables must use recalc_deliverable, which requires a build
    descriptor and always supplies its mapping.
    """
    if mapping is not None:
        mapping = TemplateMap.model_validate(mapping.model_dump())
        if set(expected) != {entry.address for entry in mapping.entries}:
            raise ValueError("Parity requires complete mapped expectation coverage")
        validate_mapped_workbook(path, mapping, gates=gates)
    if not expected:
        raise ValueError("Parity requires a nonempty authoritative expected-value map")
    workbook = load_workbook(path, data_only=True)
    failures: list[str] = []
    compared = 0
    try:
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e":
                        failures.append(
                            f"{sheet.title}!{cell.coordinate}: Excel error {cell.value}"
                        )
        for address, reference in expected.items():
            if not isinstance(reference, Decimal) or not reference.is_finite():
                raise ValueError(f"{address}: expected value must be a finite stored Decimal")
            if "!" not in address:
                raise ValueError("Expected addresses must be sheet!cell")
            sheet_name, coordinate = address.rsplit("!", 1)
            value = (
                workbook[sheet_name][coordinate].value
                if sheet_name in workbook.sheetnames
                else None
            )
            if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
                failures.append(f"{address}: missing or nonnumeric recalculated cache ({value})")
                continue
            actual = Decimal(str(value))
            if not actual.is_finite():
                failures.append(f"{address}: nonfinite recalculated cache")
                continue
            # Exact rational differences are independent of caller Decimal context.
            difference = abs(Fraction(actual) - Fraction(reference))
            limit = max(
                Fraction(gates.parity_abs), Fraction(gates.parity_rel) * abs(Fraction(reference))
            )
            compared += 1
            if difference > limit:
                failures.append(
                    f"{address}: Excel {actual} != stored {reference}; exceeds configured tolerance"
                )
        return GateResult(
            passed=not failures, failures=failures, metrics={"compared_cells": Decimal(compared)}
        )
    finally:
        workbook.close()
