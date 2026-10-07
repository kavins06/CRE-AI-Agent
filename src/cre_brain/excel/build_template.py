"""Reproducible five-year, unlevered multifamily operating mirror.

Run: python -m cre_brain.excel.build_template --output templates --config config
Taxes/value-add/debt/returns and circular solutions are separate finance modules;
this reference mirrors build_proforma without optional schedules. Firm extensions
must supply an explicit map and labelled stored calculation inputs.
"""

import argparse
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from cre_brain.config import load
from cre_brain.config.settings import GateSettings
from cre_brain.excel.models import CellMapping, TemplateMap
from cre_brain.excel.validation import validate_workbook

INPUTS = (
    "base_monthly_revenue",
    "base_monthly_operating_expenses",
    "annual_revenue_growth",
    "annual_expense_growth",
    "vacancy_rate",
    "credit_loss_rate",
    "monthly_reserves",
)
MONTHLY = (
    "gross_potential_revenue",
    "vacancy_loss",
    "credit_loss",
    "effective_gross_income",
    "operating_expenses",
    "reserves",
    "noi",
    "cash_flow",
)
ANNUAL = (
    "gross_potential_revenue",
    "effective_gross_income",
    "operating_expenses",
    "reserves",
    "noi",
    "cash_flow",
)


def build_template(directory: Path, *, gates: GateSettings) -> tuple[Path, TemplateMap]:
    directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    inputs = workbook.create_sheet("Inputs")
    monthly = workbook.create_sheet("Monthly")
    annual = workbook.create_sheet("Annual")
    entries: list[CellMapping] = []

    def add(entry: CellMapping) -> None:
        entries.append(entry)
        workbook.defined_names.add(
            DefinedName(
                entry.name,
                attr_text=f"{quote_sheetname(entry.sheet)}!{entry.cell}",
            )
        )
        if entry.formula:
            workbook[entry.sheet][entry.cell] = entry.formula
        workbook[entry.sheet][entry.cell].number_format = "0.000000"

    inputs.append(["Stored fact key", "Value", "Unit"])
    for row, key in enumerate(INPUTS, 2):
        unit = "ratio" if "rate" in key or "growth" in key else "USD/month"
        inputs.cell(row, 1, key)
        inputs.cell(row, 3, unit)
        add(
            CellMapping(
                name=key,
                sheet="Inputs",
                cell=f"B{row}",
                source="fact",
                role="input",
                key=key,
                unit=unit,
            )
        )
    monthly.append(["Period", *MONTHLY])
    for month in range(1, 61):
        row = month + 1
        year = (month - 1) // 12 + 1
        # Calendar indexes are template structure, not claimed deal facts.
        monthly.cell(row, 1, f"Month {month}")
        formulas = (
            f"=base_monthly_revenue*(1+annual_revenue_growth)^{year - 1}",
            f"=B{row}*vacancy_rate",
            f"=(B{row}-C{row})*credit_loss_rate",
            f"=B{row}-C{row}-D{row}",
            f"=base_monthly_operating_expenses*(1+annual_expense_growth)^{year - 1}",
            "=monthly_reserves",
            f"=E{row}-F{row}",
            f"=H{row}-G{row}",
        )
        for column, (key, formula) in enumerate(zip(MONTHLY, formulas, strict=True), 2):
            cell = monthly.cell(row, column).coordinate
            add(
                CellMapping(
                    name=f"month_{month}_{key}",
                    sheet="Monthly",
                    cell=cell,
                    source="calc",
                    role="output",
                    key=f"month:{month}:{key}",
                    calculation="proforma",
                    function="build_proforma",
                    unit="USD",
                    formula=formula,
                )
            )
    annual.append(["Metric", *[f"Year {year}" for year in range(1, 6)]])
    # Keep an empty B column as a label/provenance margin; values are C:G.
    annual.insert_cols(2)
    for row, key in enumerate(ANNUAL, 2):
        annual.cell(row, 1, key)
        monthly_col = chr(66 + MONTHLY.index(key))
        for year in range(1, 6):
            cell = annual.cell(row, year + 2).coordinate
            start = (year - 1) * 12 + 2
            formula = f"=SUM(Monthly!{monthly_col}{start}:{monthly_col}{start + 11})"
            add(
                CellMapping(
                    name=f"year_{year}_{key}",
                    sheet="Annual",
                    cell=cell,
                    source="calc",
                    role="output",
                    key=f"year:{year}:{key}",
                    calculation="proforma",
                    function="build_proforma",
                    unit="USD",
                    formula=formula,
                )
            )
    mapping = TemplateMap(entries=entries)
    workbook.properties.created = workbook.properties.modified = datetime(2000, 1, 1)
    for sheet in workbook:
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 38
    path = directory / "mf_standard.xlsx"
    workbook.save(path)
    # Normalize ZIP metadata as well as workbook properties for identical bytes.
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    core = ET.fromstring(members["docProps/core.xml"])
    for node in core:
        if node.tag.rsplit("}", 1)[-1] in {"created", "modified"}:
            node.text = "2000-01-01T00:00:00Z"
    members["docProps/core.xml"] = ET.tostring(core)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(members.items()):
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)
    validate_workbook(path, gates=gates)
    path.with_suffix(".map.json").write_text(mapping.model_dump_json(indent=2) + "\n")
    return path, mapping


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("templates"))
    parser.add_argument("--config", type=Path, default=Path("config"))
    args = parser.parse_args()
    path, _ = build_template(args.output, gates=load(args.config).gates)
    print(path)


if __name__ == "__main__":
    main()
