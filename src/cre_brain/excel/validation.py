"""Fail-closed supported-feature scan for trusted synthetic workbooks.

This is feature validation, not a native Office parsing sandbox. Seller uploads
must remain behind the future contained onboarding/extraction boundary.
"""

import re
import xml.etree.ElementTree as ET
import zipfile
from graphlib import CycleError, TopologicalSorter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook

from cre_brain.config.settings import GateSettings
from cre_brain.excel.models import TemplateMap

MAX_CELLS = 20000
MAX_BYTES = 20_000_000


def _native_calc_metadata(name: str, root: ET.Element, node: ET.Element) -> bool:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    if name != "xl/workbook.xml" or root.tag != f"{{{namespace}}}workbook" or node not in root:
        return False
    if len(root.findall(f"{{{namespace}}}extLst")) != 1:
        return False
    elements = list(node.iter())
    expected = [
        (f"{{{namespace}}}extLst", {}, 1),
        (f"{{{namespace}}}ext", {"uri": "{7626C862-2A13-11E5-B345-FEFF819CDC9F}"}, 1),
        ("{http://schemas.libreoffice.org/}extCalcPr", {"stringRefSyntax": "ExcelA1"}, 0),
    ]
    return [(e.tag, e.attrib, len(e)) for e in elements] == expected and all(
        not (e.text or "").strip() and not (e.tail or "").strip() for e in elements
    )


def formula_signature(formula: str) -> tuple[tuple[str, str, str], ...]:
    """Only normalize case, sheet quoting and fixed-cell absolute markers.

    These are equivalent for a saved formula at a fixed address. Operators,
    constants, functions, range boundaries and operand order remain exact.
    Named references are never substituted by dependency sets.
    """
    signature = []
    for token in Tokenizer(formula).items:
        value = token.value
        if token.type == "FUNC":
            value = value.upper()
        elif token.type == "OPERAND" and token.subtype == "RANGE":
            if "!" in value:
                sheet, cell = value.rsplit("!", 1)
                value = sheet.strip("'").replace("''", "'") + "!" + cell.replace("$", "")
            else:
                value = value.replace("$", "")
            value = value.casefold()
        signature.append((token.type, token.subtype, value))
    return tuple(signature)


def validate_mapped_workbook(path: Path, mapping: TemplateMap, *, gates: GateSettings) -> None:
    """Require the validated map's complete names, roles and formula structure."""
    formulas = validate_workbook(path, gates=gates)
    workbook = load_workbook(path)
    try:
        names = {name.casefold(): definition for name, definition in workbook.defined_names.items()}
        if set(names) != {entry.name.casefold() for entry in mapping.entries}:
            raise ValueError("Template names and map disagree")
        for entry in mapping.entries:
            destinations = list(names[entry.name.casefold()].destinations)
            if [(s.casefold(), c.replace("$", "").upper()) for s, c in destinations] != [
                (entry.sheet.casefold(), entry.cell.replace("$", ""))
            ]:
                raise ValueError(f"Map/name mismatch for {entry.name}")
            actual = formulas.get(entry.address)
            if entry.formula is not None:
                if actual is None or formula_signature(actual) != formula_signature(entry.formula):
                    raise ValueError(f"Mapped formula changed at {entry.address}")
            elif actual is not None:
                raise ValueError(f"Mapped input unexpectedly contains a formula at {entry.address}")
    finally:
        workbook.close()


def _references(value: str, sheet: str, workbook: Workbook) -> set[str]:
    if "[" in value or "]" in value or "#" in value:
        raise ValueError(f"Unsupported external/structured reference {value}")
    names = {name.casefold(): definition for name, definition in workbook.defined_names.items()}
    if value.casefold() in names:
        definition = names[value.casefold()]
        destinations = list(definition.destinations)
        if len(destinations) != 1:
            raise ValueError("Named ranges must have a single internal destination")
        sheet, value = destinations[0]
    elif "!" in value:
        sheet, value = value.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
    sheets = {name.casefold(): name for name in workbook.sheetnames}
    sheet = sheets.get(sheet.casefold(), sheet)
    value = value.upper()
    if sheet not in workbook.sheetnames or not re.fullmatch(
        r"\$?[A-Z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Z]{1,3}\$?[1-9][0-9]*)?", value
    ):
        raise ValueError(f"Unsupported or unknown reference {sheet}!{value}")
    left, top, right, bottom = range_boundaries(value)
    if left is None or top is None or right is None or bottom is None:
        raise ValueError("Whole-column/row references are unsupported")
    if right > 16384 or bottom > 1048576 or (right - left + 1) * (bottom - top + 1) > MAX_CELLS:
        raise ValueError("Reference exceeds supported cell budget")
    return {
        f"{sheet}!{workbook[sheet].cell(row, col).coordinate}"
        for row in range(top, bottom + 1)
        for col in range(left, right + 1)
    }


def validate_workbook(path: Path, *, gates: GateSettings) -> dict[str, str]:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only macro-free .xlsx trusted templates are supported")
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 2000 or sum(m.file_size for m in members) > MAX_BYTES:
            raise ValueError("Workbook exceeds supported package budget")
        for member in members:
            name = member.filename.lower()
            if any(
                part in name
                for part in (
                    "vba",
                    "externallink",
                    "embedding",
                    "activex",
                    "pivot",
                    "connections",
                    "querytable",
                )
            ):
                raise ValueError(f"Unsupported workbook feature: {member.filename}")
            if name.endswith(".rels"):
                root = ET.fromstring(archive.read(member))
                if any(e.get("TargetMode") == "External" for e in root):
                    raise ValueError("External links are unsupported")
                if any(
                    any(
                        feature in e.get("Type", "").casefold() for feature in ("vba", "macrosheet")
                    )
                    for e in root
                ):
                    raise ValueError("Macro/VBA relationships are unsupported")
            if name.endswith(".xml"):
                root = ET.fromstring(archive.read(member))
                for node in root.iter():
                    if any(
                        feature in node.get("ContentType", "").casefold()
                        for feature in ("macroenabled", "vba", "macrosheet")
                    ):
                        raise ValueError("Macro-enabled/VBA content types are unsupported")
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag == "extLst" and not _native_calc_metadata(member.filename, root, node):
                        raise ValueError(f"Unsupported workbook feature: {tag}")
                    if tag in {"oleObject", "dataTable", "tableParts"}:
                        raise ValueError(f"Unsupported workbook feature: {tag}")
                    if tag == "f" and node.get("t", "normal") != "normal":
                        raise ValueError("Array/shared/data-table formulas are unsupported")
    workbook = load_workbook(path, keep_links=True)
    try:
        if any(sheet.defined_names for sheet in workbook):
            raise ValueError("Worksheet-local defined names are unsupported")
        if len({name.casefold() for name in workbook.defined_names}) != len(workbook.defined_names):
            raise ValueError("Ambiguous case-insensitive defined names")
        if workbook.calculation and workbook.calculation.iterate:
            raise ValueError("Iterative calculation is unsupported")
        for name in workbook.defined_names:
            _references(name, workbook.sheetnames[0], workbook)
        formulas: dict[str, str] = {}
        graph: dict[str, set[str]] = {}
        count = 0
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise ValueError("Worksheet exceeds supported cell budget")
            for row in sheet:
                for cell in row:
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("Workbook exceeds supported cell budget")
                    if cell.data_type != "f":
                        continue
                    if not isinstance(cell.value, str) or len(cell.value) > 4096:
                        raise ValueError("Unsupported formula")
                    address = f"{sheet.title}!{cell.coordinate}"
                    formulas[address] = cell.value
                    dependencies: set[str] = set()
                    for token in Tokenizer(cell.value).items:
                        if token.type == "FUNC" and token.subtype == "OPEN":
                            function = token.value[:-1].upper()
                            if function not in gates.excel_functions:
                                raise ValueError(f"{address}: unsupported function {function}")
                        elif token.type == "OPERAND" and token.subtype == "RANGE":
                            dependencies.update(_references(token.value, sheet.title, workbook))
                        elif token.type == "OPERAND" and token.subtype == "ERROR":
                            raise ValueError(f"{address}: formula contains an invalid reference")
                        elif token.type in {"ARRAY"}:
                            raise ValueError(f"{address}: array formulas are unsupported")
                    graph[address] = dependencies
        try:
            tuple(TopologicalSorter(graph).static_order())
        except CycleError as exc:
            raise ValueError("Circular formula references are unsupported") from exc
        return formulas
    finally:
        workbook.close()
