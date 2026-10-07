"""Immutable source observations, never verified facts or finance calculations."""

from __future__ import annotations

import hashlib
import json
from typing import Annotated, Literal, Self

from openpyxl.utils.cell import coordinate_to_tuple
from pydantic import BaseModel, ConfigDict, Field, StringConstraints, model_validator

from cre_brain.domain.base import TenantScope
from cre_brain.domain.models import Provenance

DocId = Annotated[str, StringConstraints(pattern=r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$")]
Digest = Annotated[str, StringConstraints(pattern=r"^[0-9a-f]{64}$")]
Text = Annotated[str, StringConstraints(max_length=1_000_000)]
SheetName = Annotated[str, StringConstraints(min_length=1, max_length=128)]
CellAddress = Annotated[str, StringConstraints(pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")]
Version = Annotated[str, StringConstraints(min_length=1, max_length=128)]
Coordinate = Annotated[float, Field(ge=0, le=14400, allow_inf_nan=False)]
Dimension = Annotated[float, Field(gt=0, le=14400, allow_inf_nan=False)]
type Box = tuple[Coordinate, Coordinate, Coordinate, Coordinate]


class PreparseError(ValueError):
    """Reject malformed, unsupported, unconfined or over-budget input, without truncation."""


class ParserUnavailable(PreparseError):
    """An optional parser cannot run; this is not successful empty parsing."""


class Boundary(BaseModel):
    model_config = ConfigDict(
        extra="forbid",
        frozen=True,
        strict=True,
        validate_default=True,
        revalidate_instances="always",
        allow_inf_nan=False,
    )


class Limits(Boundary):
    max_file_bytes: Annotated[int, Field(ge=1, le=64 * 1024 * 1024)] = 16 * 1024 * 1024
    max_zip_members: Annotated[int, Field(ge=1, le=1024)] = 128
    max_zip_member_bytes: Annotated[int, Field(ge=1, le=16 * 1024 * 1024)] = 8 * 1024 * 1024
    max_zip_expanded_bytes: Annotated[int, Field(ge=1, le=64 * 1024 * 1024)] = 32 * 1024 * 1024
    max_zip_ratio: Annotated[int, Field(ge=1, le=1000)] = 200
    max_cells: Annotated[int, Field(ge=1, le=100_000)] = 20_000
    max_rows: Annotated[int, Field(ge=1, le=100_000)] = 10_000
    max_columns: Annotated[int, Field(ge=1, le=1024)] = 256
    max_sheets: Annotated[int, Field(ge=1, le=128)] = 32
    max_pages: Annotated[int, Field(ge=1, le=200)] = 100
    max_cell_chars: Annotated[int, Field(ge=1, le=1_000_000)] = 32_768
    max_text_chars: Annotated[int, Field(ge=1, le=2_000_000)] = 1_000_000
    max_output_bytes: Annotated[int, Field(ge=1, le=32 * 1024 * 1024)] = 8 * 1024 * 1024
    worker_seconds: Annotated[int, Field(ge=1, le=120)] = 60
    worker_cpu_seconds: Annotated[int, Field(ge=1, le=60)] = 30
    worker_memory_mb: Annotated[int, Field(ge=256, le=4096)] = 2048


class CellAnchor(Boundary):
    doc_id: DocId
    sheet: SheetName
    cell: CellAddress


class PageAnchor(Boundary):
    doc_id: DocId
    page: Annotated[int, Field(ge=1, le=200)]
    bbox: Box
    coord_origin: Literal["TOPLEFT"] = "TOPLEFT"

    @model_validator(mode="after")
    def ordered(self) -> Self:
        if self.bbox[0] > self.bbox[2] or self.bbox[1] > self.bbox[3]:
            raise ValueError("Bounding box must be ordered left/top/right/bottom")
        return self


class ParsedCell(Boundary):
    cell_id: Digest
    anchor: CellAnchor
    kind: Literal["text", "number", "boolean", "formula", "error", "date", "empty"]
    text: Text
    cached_text: Text | None = None
    provenance: Provenance

    @model_validator(mode="after")
    def source_matches(self) -> Self:
        expected = Provenance(
            doc_id=self.anchor.doc_id,
            sheet=self.anchor.sheet,
            cell=self.anchor.cell,
            quote=self.text,
        )
        if self.provenance != expected:
            raise ValueError("Provenance must match exact cell observation")
        return self


class ParsedTable(Boundary):
    table_id: Digest
    sheet: SheetName
    cells: Annotated[tuple[ParsedCell, ...], Field(max_length=100_000)]


class ParsedText(Boundary):
    text_id: Digest
    anchor: PageAnchor
    text: Text
    provenance: Provenance

    @model_validator(mode="after")
    def source_matches(self) -> Self:
        expected = Provenance(
            doc_id=self.anchor.doc_id,
            page=self.anchor.page,
            bbox=self.anchor.bbox,
            quote=self.text,
        )
        if self.provenance != expected:
            raise ValueError("Provenance must match exact native text observation")
        return self


class ParsedPage(Boundary):
    page: Annotated[int, Field(ge=1, le=200)]
    width: Dimension
    height: Dimension


class PdfText(Boundary):
    page: Annotated[int, Field(ge=1, le=200)]
    bbox: Box
    text: Text


class PdfContent(Boundary):
    pages: Annotated[tuple[ParsedPage, ...], Field(max_length=200)]
    texts: Annotated[tuple[PdfText, ...], Field(max_length=100_000)]
    parser_versions: Annotated[tuple[tuple[Version, Version], ...], Field(max_length=16)]


class ParsedDocument(Boundary):
    schema_version: Literal["1.0.0"] = "1.0.0"
    doc_id: DocId
    scope: TenantScope
    source_name: Annotated[str, StringConstraints(min_length=1, max_length=255)]
    source_sha256: Digest
    source_bytes: Annotated[int, Field(ge=0, le=64 * 1024 * 1024)]
    parser: Literal["native-csv", "native-xlsx", "docling-native", "reducto"]
    parser_versions: Annotated[tuple[tuple[Version, Version], ...], Field(max_length=16)]
    config_sha256: Digest
    limits: Limits
    status: Literal["parsed", "needs_review", "empty"]
    tables: Annotated[tuple[ParsedTable, ...], Field(max_length=128)] = ()
    texts: Annotated[tuple[ParsedText, ...], Field(max_length=100_000)] = ()
    pages: Annotated[tuple[ParsedPage, ...], Field(max_length=200)] = ()
    pages_without_text: Annotated[tuple[int, ...], Field(max_length=200)] = ()
    warnings: Annotated[
        tuple[
            Literal[
                "external_links_not_followed",
                "hidden_sheets_included",
                "native_pdf_no_ocr_layout_or_table_inference",
                "native_pdf_no_visibility_verification",
                "pages_without_native_text",
            ],
            ...,
        ],
        Field(max_length=8),
    ] = ()

    @model_validator(mode="after")
    def coherent(self) -> Self:
        if (
            self.source_name in (".", "..")
            or "/" in self.source_name
            or "\\" in self.source_name
            or "\0" in self.source_name
        ):
            raise ValueError("Source name must not contain a path")
        cells = [c for table in self.tables for c in table.cells]
        if (
            len(self.tables) > self.limits.max_sheets
            or len(cells) + len(self.texts) > self.limits.max_cells
            or len(self.pages) > self.limits.max_pages
            or self.source_bytes > self.limits.max_file_bytes
        ):
            raise ValueError("Parsed document exceeds resource limits")
        strings = [c.text for c in cells] + [t.text for t in self.texts]
        strings += [c.cached_text for c in cells if c.cached_text is not None]
        if any(len(s) > self.limits.max_cell_chars for s in strings) or (
            sum(map(len, strings)) > self.limits.max_text_chars
        ):
            raise ValueError("Parsed text exceeds resource limits")
        anchors = [c.anchor.doc_id for c in cells] + [t.anchor.doc_id for t in self.texts]
        if any(doc_id != self.doc_id for doc_id in anchors):
            raise ValueError("Anchor belongs to another document")
        ids = [c.cell_id for c in cells] + [t.text_id for t in self.texts]
        if len(set(ids)) != len(ids):
            raise ValueError("Duplicate observation ID")
        for table in self.tables:
            addresses = [c.anchor.cell for c in table.cells]
            if len(set(addresses)) != len(addresses) or any(
                c.anchor.sheet != table.sheet for c in table.cells
            ):
                raise ValueError("Duplicate or mismatched cell anchors")
            for address in addresses:
                row, column = coordinate_to_tuple(address)
                if row > self.limits.max_rows or column > self.limits.max_columns:
                    raise ValueError("Cell anchor exceeds row/column limits")
        if len(set(t.sheet for t in self.tables)) != len(self.tables):
            raise ValueError("Duplicate sheet names")
        if not self.parser_versions or len(set(n for n, _ in self.parser_versions)) != (
            len(self.parser_versions)
        ):
            raise ValueError("Parser versions must be present and unambiguous")
        pages = {page.page: page for page in self.pages}
        if tuple(pages) != tuple(range(1, len(self.pages) + 1)):
            raise ValueError("Page numbers must be contiguous and ordered")
        for text in self.texts:
            page = pages.get(text.anchor.page)
            if (
                page is None
                or text.anchor.bbox[2] > page.width
                or text.anchor.bbox[3] > page.height
            ):
                raise ValueError("Text anchor must be within its page")
        empty = tuple(sorted(set(pages) - {text.anchor.page for text in self.texts}))
        if self.pages_without_text != empty:
            raise ValueError("Pages without native text must be reported")
        if self.pages and self.status != ("needs_review" if empty else "parsed"):
            raise ValueError("PDF status must report pages without native text")
        if len({t.table_id for t in self.tables}) != len(self.tables):
            raise ValueError("Duplicate table ID")
        if self.parser in ("native-csv", "native-xlsx"):
            if self.pages or self.texts or self.status != ("parsed" if cells else "empty"):
                raise ValueError(
                    "Native cells must not invent PDF anchors or successful empty data"
                )
        elif not self.pages or self.tables:
            raise ValueError("PDF text parsing requires pages, not fabricated native tables")
        return self


def digest(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode()
    ).hexdigest()


class Budget:
    def __init__(self, limits: Limits):
        self.limits = limits
        self.cells = 0
        self.chars = 0

    def add(self, text: str, cached: str | None = None) -> None:
        self.cells += 1
        if self.cells > self.limits.max_cells:
            raise PreparseError("Cell/text item limit exceeded")
        for value in (text, cached):
            if value is None:
                continue
            self.chars += len(value)
            if len(value) > self.limits.max_cell_chars or self.chars > self.limits.max_text_chars:
                raise PreparseError("Text resource limit exceeded")
