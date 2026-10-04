"""Native CSV and OOXML observations: exact decoded lexemes, no formula execution."""

from __future__ import annotations

import csv
import io
import posixpath
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from cre_brain.domain.models import Provenance
from cre_brain.extraction.preparse.models import (
    Budget,
    CellAnchor,
    Limits,
    ParsedCell,
    ParsedTable,
    PreparseError,
    digest,
)

MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _cell(
    seed: str,
    doc_id: str,
    sheet: str,
    address: str,
    kind: str,
    text: str,
    budget: Budget,
    cached: str | None = None,
) -> ParsedCell:
    budget.add(text, cached)
    return ParsedCell.model_validate(
        {
            "cell_id": digest((seed, sheet, address, kind, text, cached)),
            "anchor": CellAnchor(doc_id=doc_id, sheet=sheet, cell=address),
            "kind": kind,
            "text": text,
            "cached_text": cached,
            "provenance": Provenance(doc_id=doc_id, sheet=sheet, cell=address, quote=text),
        }
    )


def csv_table(raw: bytes, doc_id: str, seed: str, limits: Limits, delimiter: str) -> ParsedTable:
    budget = Budget(limits)
    cells = []
    try:
        decoded = raw.decode("utf-8-sig")
        _validate_csv(decoded, delimiter, limits)
        reader = csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter, strict=True)
        for row_number, row in enumerate(reader, 1):
            if row_number > limits.max_rows or len(row) > limits.max_columns:
                raise PreparseError("CSV row/column limit exceeded")
            for column, text in enumerate(row, 1):
                cells.append(
                    _cell(
                        seed,
                        doc_id,
                        "CSV",
                        f"{get_column_letter(column)}{row_number}",
                        "text",
                        text,
                        budget,
                    )
                )
    except (UnicodeError, csv.Error) as error:
        raise PreparseError(
            "CSV must be valid UTF-8 with valid explicit-dialect quoting"
        ) from error
    return ParsedTable(table_id=digest((seed, "CSV")), sheet="CSV", cells=tuple(cells))


def _validate_csv(content: str, delimiter: str, limits: Limits) -> None:
    # Bound fields before csv.reader materializes a row; reject quotes in bare fields.
    state = "start"
    row, column, chars = 1, 1, 0
    previous_cr = False
    for character in content:
        if previous_cr and character == "\n":
            previous_cr = False
            continue
        previous_cr = False
        if row > limits.max_rows:
            raise PreparseError("CSV row limit exceeded")
        if state == "quoted":
            if character == '"':
                state = "closed"
            else:
                chars += 1
        elif state == "closed" and character == '"':
            state = "quoted"
            chars += 1
        elif character == delimiter:
            column += 1
            state, chars = "start", 0
            if column > limits.max_columns:
                raise PreparseError("CSV column limit exceeded")
        elif character in ("\r", "\n"):
            row += 1
            column, chars, state = 1, 0, "start"
            previous_cr = character == "\r"
        elif state == "start" and character == '"':
            state = "quoted"
        elif state == "closed" or character == '"':
            raise PreparseError("Malformed CSV quoting")
        else:
            state = "bare"
            chars += 1
        if chars > limits.max_cell_chars:
            raise PreparseError("CSV cell text limit exceeded")
    if state == "quoted":
        raise PreparseError("Unclosed CSV quote")


def _xml(raw: bytes) -> ET.Element:
    inspected = raw.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in inspected or b"<!ENTITY" in inspected:
        raise PreparseError("XML entities and DTDs are not permitted")
    try:
        return ET.fromstring(raw)
    except ET.ParseError as error:
        raise PreparseError("Malformed spreadsheet XML") from error


def _zip_members(raw: bytes, limits: Limits) -> tuple[dict[str, bytes], bool]:
    members: dict[str, bytes] = {}
    total = 0
    external = False
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_zip_members:
                raise PreparseError("ZIP member count exceeded")
            for info in infos:
                parts = info.filename.split("/")
                if (
                    info.filename.startswith("/")
                    or "\\" in info.filename
                    or "\0" in info.filename
                    or info.orig_filename != info.filename
                    or ":" in info.filename
                    or any(part in (".", "..") for part in parts)
                    or info.filename in members
                    or info.flag_bits & 1
                    or info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)
                    or (info.external_attr >> 16) & 0o170000 == 0o120000
                ):
                    raise PreparseError("Unsafe, duplicate, encrypted or unsupported ZIP member")
                if "vbaproject" in info.filename.lower():
                    raise PreparseError("Macro-bearing spreadsheets are unsupported")
                if info.is_dir():
                    continue
                total += info.file_size
                if (
                    info.file_size > limits.max_zip_member_bytes
                    or total > limits.max_zip_expanded_bytes
                    or info.file_size > max(1, info.compress_size) * limits.max_zip_ratio
                ):
                    raise PreparseError("ZIP expansion/ratio limit exceeded")
                with archive.open(info) as member:
                    content = member.read(limits.max_zip_member_bytes + 1)
                if len(content) != info.file_size or len(content) > limits.max_zip_member_bytes:
                    raise PreparseError("ZIP member byte limit exceeded")
                if info.filename.endswith((".xml", ".rels")):
                    root = _xml(content)
                    if any(element.get("TargetMode") == "External" for element in root.iter()):
                        external = True
                    if b"macroEnabled" in content:
                        raise PreparseError("Macro-enabled spreadsheet content is unsupported")
                members[info.filename] = content
    except (zipfile.BadZipFile, OSError, RuntimeError, EOFError) as error:
        raise PreparseError("Invalid spreadsheet ZIP archive") from error
    return members, external


def _strings(root: ET.Element) -> str:
    container = root.find(f"{MAIN}is") if root.tag == f"{MAIN}c" else root
    if container is None or container.tag not in (f"{MAIN}is", f"{MAIN}si"):
        raise PreparseError("Invalid spreadsheet string cell")
    strings = []
    for child in container:
        node = child.find(f"{MAIN}t") if child.tag == f"{MAIN}r" else child
        if node is not None and node.tag == f"{MAIN}t":
            strings.append(node.text or "")
    return "".join(strings)


def xlsx_tables(
    raw: bytes,
    doc_id: str,
    seed: str,
    limits: Limits,
) -> tuple[tuple[ParsedTable, ...], bool, bool]:
    members, external = _zip_members(raw, limits)
    budget = Budget(limits)
    try:
        workbook = _xml(members["xl/workbook.xml"])
        relationships = _xml(members["xl/_rels/workbook.xml.rels"])
        targets = {}
        for relationship in relationships:
            if relationship.get("TargetMode") == "External":
                continue
            target = relationship.get("Target", "")
            part = posixpath.normpath(
                target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            )
            if part.startswith("../") or ":" in part or "\\" in part:
                raise PreparseError("Spreadsheet relationship escapes package")
            targets[relationship.attrib["Id"]] = part
        shared = []
        if "xl/sharedStrings.xml" in members:
            shared = [_strings(si) for si in _xml(members["xl/sharedStrings.xml"])]
            if len(shared) > limits.max_cells or sum(map(len, shared)) > limits.max_text_chars:
                raise PreparseError("Shared string resource limit exceeded")
        sheets = list(workbook.iter(f"{MAIN}sheet"))
        if workbook.tag != f"{MAIN}workbook" or not sheets:
            raise PreparseError("Unsupported spreadsheet namespace or no worksheets")
        if len(sheets) > limits.max_sheets:
            raise PreparseError("Sheet count exceeded")
        hidden = any(sheet.get("state", "visible") != "visible" for sheet in sheets)
        # Native reader validates workbook structure. Do not iterate declared dimensions or
        # use its float/date coercions for source observations; read bounded XML lexemes.
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if book.sheetnames != [sheet.attrib["name"] for sheet in sheets]:
                raise PreparseError("Workbook sheet metadata mismatch")
        finally:
            book.close()
        tables = []
        for sheet in sheets:
            name = sheet.attrib["name"]
            part = targets[sheet.attrib[f"{REL}id"]]
            cells = []
            worksheet = _xml(members[part])
            if worksheet.tag != f"{MAIN}worksheet":
                raise PreparseError("Only native worksheet cells are supported")
            data_sections = worksheet.findall(f"{MAIN}sheetData")
            if len(data_sections) != 1:
                raise PreparseError("Worksheet must contain exactly one sheet data section")
            sheet_data = data_sections[0]
            for node in (
                cell for row in sheet_data.findall(f"{MAIN}row") for cell in row.findall(f"{MAIN}c")
            ):
                address = node.attrib["r"]
                row, column = coordinate_to_tuple(address)
                if row > limits.max_rows or column > limits.max_columns:
                    raise PreparseError("Spreadsheet row/column limit exceeded")
                value = node.find(f"{MAIN}v")
                text = value.text if value is not None and value.text is not None else ""
                cached = None
                kind = node.get("t", "n")
                formula = node.find(f"{MAIN}f")
                if formula is not None:
                    if formula.attrib or formula.text is None:
                        raise PreparseError("Shared/array/data-table formulas are unsupported")
                    cached = value.text if value is not None else None
                    text, kind = "=" + formula.text, "formula"
                elif kind == "s":
                    index = int(text)
                    if index < 0 or index >= len(shared):
                        raise PreparseError("Invalid shared-string index")
                    text, kind = shared[index], "text"
                elif kind == "inlineStr":
                    text, kind = _strings(node), "text"
                elif kind == "str":
                    kind = "text"
                elif kind == "n":
                    kind = "number" if text else "empty"
                elif kind == "b":
                    if text not in ("0", "1"):
                        raise PreparseError("Invalid boolean cell")
                    kind = "boolean"
                elif kind == "e":
                    kind = "error"
                elif kind == "d":
                    kind = "date"
                else:
                    raise PreparseError("Unsupported cell type")
                cells.append(_cell(seed, doc_id, name, address, kind, text, budget, cached))
            tables.append(
                ParsedTable(table_id=digest((seed, name)), sheet=name, cells=tuple(cells))
            )
        return tuple(tables), external, hidden
    except (KeyError, ValueError, IndexError, TypeError) as error:
        if isinstance(error, PreparseError):
            raise
        raise PreparseError("Malformed or unsupported native spreadsheet") from None
